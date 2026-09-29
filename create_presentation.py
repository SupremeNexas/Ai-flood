"""
Script to create a professional, research-grade PowerPoint presentation
for the AI Flood Detection Minor Project.
Uses python-pptx to generate AI_Flood_Detection_Minor_Project.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Color Palette Constants ---
C_BG = RGBColor(11, 19, 43)          # #0B132B Deep Navy
C_CARD = RGBColor(30, 41, 59)        # #1E293B Slate Card Fill
C_CARD_DARK = RGBColor(15, 23, 42)   # #0F172A Dark Slate Card
C_BORDER = RGBColor(51, 65, 85)      # #334155 Slate Border
C_BORDER_CYAN = RGBColor(2, 132, 199)# #0284C7 Accent Border
C_CYAN = RGBColor(0, 229, 255)       # #00E5FF Bright Cyan
C_CYAN_LIGHT = RGBColor(56, 189, 248)# #38BDF8 Sky Cyan
C_GREEN = RGBColor(16, 185, 129)     # #10B981 Emerald Green
C_AMBER = RGBColor(245, 158, 11)     # #F59E0B Amber / Warning
C_WHITE = RGBColor(255, 255, 255)    # #FFFFFF White
C_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Muted Gray
C_TEXT_LIGHT = RGBColor(241, 245, 249) # #F1F5F9 Off-white

FONT_HEADING = "Arial"
FONT_BODY = "Arial"

def set_slide_background(slide, color):
    """Fills slide background with a solid color."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, badge_text, title_text, subtitle_text=""):
    """Adds a standardized top header zone to the slide."""
    # Category badge
    if badge_text:
        badge_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), Inches(0.35), Inches(3.2), Inches(0.3)
        )
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = C_BORDER_CYAN
        badge_box.line.fill.background()
        tf = badge_box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = badge_text.upper()
        p.font.name = FONT_HEADING
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.7), Inches(12.13), Inches(0.55)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_HEADING
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Subtitle
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.25), Inches(12.13), Inches(0.35)
        )
        tf = sub_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = subtitle_text
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_CYAN_LIGHT

def add_footer(slide, slide_num, total_slides=16):
    """Adds standardized footer to the slide."""
    footer_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(7.05), Inches(10.0), Inches(0.3)
    )
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = "AI-Based Flood Detection & Mapping  |  Sentinel-1 SAR • PyTorch U-Net  |  Academic Minor Project"
    p.font.name = FONT_BODY
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED

    num_box = slide.shapes.add_textbox(
        Inches(11.5), Inches(7.05), Inches(1.2), Inches(0.3)
    )
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"{slide_num} / {total_slides}"
    p.font.name = FONT_BODY
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED
    p.alignment = PP_ALIGN.RIGHT

def create_card(slide, left, top, width, height, bg_color=C_CARD, border_color=C_BORDER, line_width=1):
    """Creates a stylized card container."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(line_width)
    else:
        card.line.fill.background()
    return card

def add_card_header(slide, left, top, width, text, color=C_CYAN, size=13):
    """Adds a bold header inside a card."""
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEADING
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return tb

def add_speaker_notes(slide, notes_text):
    """Sets speaker notes for a slide."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text.strip()

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    assets_dir = "/Users/supryo/Desktop/Ai-flood"

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, C_BG)

    # Outer decorative card
    create_card(slide1, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3), bg_color=C_CARD_DARK, border_color=C_BORDER_CYAN, line_width=1.5)

    # Title badge
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(3.8), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_BORDER_CYAN
    badge.line.fill.background()
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "ACADEMIC MINOR PROJECT • AI & REMOTE SENSING"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    tb_title = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(6.0), Inches(1.5))
    tf = tb_title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AI-Based Flood Detection & Mapping Using Satellite Imagery"
    p.font.name = FONT_HEADING
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Subtitle
    tb_sub = slide1.shapes.add_textbox(Inches(1.0), Inches(2.95), Inches(6.0), Inches(0.5))
    tf = tb_sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Sen1Floods11 • Sentinel-1 SAR Dual-Polarization • PyTorch U-Net"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = C_CYAN_LIGHT

    # Highlights box on left
    create_card(slide1, Inches(1.0), Inches(3.5), Inches(5.8), Inches(2.4), bg_color=C_CARD, border_color=C_BORDER)
    tb_hl = slide1.shapes.add_textbox(Inches(1.15), Inches(3.6), Inches(5.5), Inches(2.2))
    tf = tb_hl.text_frame
    tf.word_wrap = True
    lines = [
        ("• Modality: ", "Sentinel-1 SAR C-Band (VV + VH Dual Polarization)"),
        ("• Core Model: ", "Optimized PyTorch U-Net with Group Normalization (G=8)"),
        ("• Benchmark: ", "Sen1Floods11 Global Dataset (11 Events across 6 Continents)"),
        ("• Evaluation: ", "100% Full Test Split (90/90 Chips, 20.5M Valid Pixels)"),
        ("• Performance: ", "0.5548 Test IoU  |  0.7137 Test Dice/F1  |  93.34% Accuracy"),
        ("• Application: ", "Interactive Streamlit Web Dashboard + Custom GeoTIFF Ingestion")
    ]
    for i, (bold_prefix, text_suffix) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = bold_prefix
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = C_CYAN
        r2 = p.add_run()
        r2.text = text_suffix
        r2.font.size = Pt(11)
        r2.font.color.rgb = C_TEXT_LIGHT

    # Tech stack tag
    tb_tech = slide1.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(5.8), Inches(0.35))
    tf = tb_tech.text_frame
    p = tf.paragraphs[0]
    p.text = "TECH STACK: Python 3.11  •  PyTorch  •  U-Net  •  Streamlit  •  Rasterio  •  Torchvision"
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_MUTED

    # Real visual on right
    img_path = os.path.join(assets_dir, "outputs/visualizations/final_test/1_good_prediction_USA_905409.png")
    if os.path.exists(img_path):
        slide1.shapes.add_picture(img_path, Inches(7.1), Inches(1.3), width=Inches(5.3))
        # Image caption
        tb_cap = slide1.shapes.add_textbox(Inches(7.1), Inches(6.1), Inches(5.3), Inches(0.3))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Verified Model Output: 5-Panel Flood Segmentation on USA Test Chip (IoU: 0.9350)"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(slide1, """
Good morning/afternoon respected members of the evaluation committee and project guides.
Welcome to the final presentation of our academic minor project: 'AI-Based Flood Detection and Mapping Using Satellite Imagery'.
In this project, we have designed, trained, and evaluated an end-to-end deep learning semantic segmentation system using Sentinel-1 Synthetic Aperture Radar (SAR) imagery on the global Sen1Floods11 benchmark.
Our model utilizes an optimized PyTorch U-Net with Group Normalization and a domain-specific Masked BCE-Dice loss, achieving a verified Test IoU of 0.5548 and Dice score of 0.7137 across 100% of the 90-chip official test split.
We have also deployed a live interactive Streamlit application for automated flood mapping and custom GeoTIFF ingestion. Let us dive into the problem statement.
""")

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, C_BG)
    add_header(slide2, "Background & Motivation", "Disaster Response & The Optical Cloud Barrier", "Why automated radar-based flood mapping is critical for disaster management")
    add_footer(slide2, 2)

    # 3 Column Cards
    col_w = Inches(3.8)
    card_h = Inches(5.0)

    # Card 1: Flooding Crisis
    create_card(slide2, Inches(0.6), Inches(1.75), col_w, card_h)
    add_card_header(slide2, Inches(0.8), Inches(1.9), col_w - Inches(0.4), "1. The Global Flood Crisis", color=C_AMBER, size=13)
    tb = slide2.shapes.add_textbox(Inches(0.8), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    pts1 = [
        "• Most Frequent Natural Disaster: Floods account for over 40% of all disaster events worldwide.",
        "• Rapid Inundation: Floodwaters peak within hours, isolating communities and cutting off roads.",
        "• Emergency Response Need: Disaster relief teams require immediate, accurate spatial maps of inundated zones to allocate rescue resources.",
        "• Ground Survey Limitations: Ground teams are hindered by submerged infrastructure and hazardous conditions."
    ]
    for i, pt in enumerate(pts1):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(10)

    # Card 2: Optical Sensor Failure
    create_card(slide2, Inches(4.766), Inches(1.75), col_w, card_h)
    add_card_header(slide2, Inches(4.966), Inches(1.9), col_w - Inches(0.4), "2. The Optical Sensor Failure", color=RGBColor(239, 68, 68), size=13)
    tb = slide2.shapes.add_textbox(Inches(4.966), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    pts2 = [
        "• Cloud Cover Blindness: Floods are caused by intense precipitation and cyclonic storm systems, generating 90-100% cloud cover.",
        "• Optical Satellite Block: Sensors like Sentinel-2 and Landsat operate in visible/NIR spectra and cannot penetrate cloud layers or haze.",
        "• Delayed Observation: Usable optical imagery is often unavailable until days or weeks after peak floodwater levels recede.",
        "• Manual Mapping Delay: Manual photo-interpretation of large satellite swaths is too slow for real-time response."
    ]
    for i, pt in enumerate(pts2):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(10)

    # Card 3: AI + Radar Solution
    create_card(slide2, Inches(8.933), Inches(1.75), col_w, card_h, border_color=C_BORDER_CYAN)
    add_card_header(slide2, Inches(9.133), Inches(1.9), col_w - Inches(0.4), "3. The AI + SAR Solution", color=C_CYAN, size=13)
    tb = slide2.shapes.add_textbox(Inches(9.133), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    pts3 = [
        "• All-Weather Penetration: Synthetic Aperture Radar (SAR) microwaves penetrate storm clouds, rain, and operate day and night.",
        "• Distinct Physics: Smooth water specularly reflects radar pulses away, appearing dark against rough land.",
        "• Deep Learning Segmentation: Custom U-Net architecture automates pixel-level water classification directly from dual-polarization SAR.",
        "• Sub-Second Mapping: Converts complex radar rasters into actionable flood extent maps in under 50ms per chip."
    ]
    for i, pt in enumerate(pts3):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(10)

    add_speaker_notes(slide2, """
Let us examine the core problem motivating this research.
Flooding is the most catastrophic and frequent natural disaster globally. Rapid disaster response requires accurate, near-real-time maps of inundated regions.
However, traditional optical satellites—such as Sentinel-2 or Landsat—rely on visible and near-infrared light. Because flood events are triggered by severe storms and monsoon systems, optical satellites are almost always blinded by thick cloud cover.
Ground surveys are equally limited because roads and bridges are submerged.
This is where Synthetic Aperture Radar (SAR) combined with Deep Learning becomes essential. Radar microwaves penetrate through clouds, rain, and darkness, while our deep learning pipeline automatically extracts flooded areas at the pixel level.
""")

    # =========================================================================
    # SLIDE 3: OBJECTIVES
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, C_BG)
    add_header(slide3, "Scope & Goals", "Project Objectives & Technical Deliverables", "Five core research and engineering milestones established for the system")
    add_footer(slide3, 3)

    objectives = [
        ("1. Dataset Ingestion & Zero-Leakage Audit", "Ingest and validate the global Sen1Floods11 benchmark across 11 flood events. Enforce strict mutual exclusivity across train, validation, and test splits with verified zero data leakage.", C_CYAN),
        ("2. Dual-Polarization SAR Feature Extraction", "Leverage Sentinel-1 C-band SAR backscatter in both VV (vertical) and VH (cross-polarization) channels, normalized from decibels to handle specular water reflection and land volume scattering.", C_GREEN),
        ("3. Optimized U-Net Architecture", "Implement an efficient PyTorch U-Net with Group Normalization (G=8) to ensure batch-size invariant normalization stability and spatial dropout to prevent overfitting.", C_CYAN_LIGHT),
        ("4. Masked Loss Formulation & Full Benchmark Evaluation", "Formulate a hybrid loss combining Binary Cross-Entropy and Dice loss calculated exclusively on valid pixels (ignoring -1 invalid labels). Rigorously evaluate across 100% of the 90 official test chips.", C_AMBER),
        ("5. Interactive Web Application & Geospatial Deployment", "Develop a lightweight, high-performance Streamlit dashboard supporting benchmark chip inspection, real-time threshold tuning, nominal flooded area estimation, and custom GeoTIFF upload.", C_WHITE)
    ]

    card_y = 1.75
    for i, (title, desc, accent) in enumerate(objectives):
        top_pos = Inches(card_y + i * 1.0)
        create_card(slide3, Inches(0.6), top_pos, Inches(12.133), Inches(0.9), bg_color=C_CARD, border_color=C_BORDER)

        # Pill badge with number
        p_badge = slide3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), top_pos + Inches(0.18), Inches(0.55), Inches(0.55))
        p_badge.fill.solid()
        p_badge.fill.fore_color.rgb = accent
        p_badge.line.fill.background()
        tf = p_badge.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_BG
        p.alignment = PP_ALIGN.CENTER

        # Text box
        tb = slide3.shapes.add_textbox(Inches(1.55), top_pos + Inches(0.1), Inches(11.0), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = accent

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(slide3, """
To address the research challenge, we set five clear engineering and academic objectives:
First, to ingest and audit the global Sen1Floods11 benchmark dataset, ensuring strict pairwise mutual exclusivity with zero data leakage across splits.
Second, to exploit the physics of Sentinel-1 dual-polarization SAR backscatter in both VV and VH channels.
Third, to construct an optimized PyTorch U-Net architecture using Group Normalization with 8 groups, eliminating small-batch normalization instability.
Fourth, to formulate a domain-specific Masked BCE plus Dice loss that strictly ignores invalid -1 pixels and counters the 12.5% foreground class imbalance, followed by full 90-chip test evaluation.
Fifth, to build an interactive Streamlit application with custom GeoTIFF upload and nominal flooded area estimation.
""")

    # =========================================================================
    # SLIDE 4: WHY SENTINEL-1 SAR?
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, C_BG)
    add_header(slide4, "Radar Remote Sensing", "Why Sentinel-1 Synthetic Aperture Radar (SAR)?", "Physical principles of microwave backscatter and dual-polarization discrimination")
    add_footer(slide4, 4)

    # Left Column: Physics & Channels (6.0 inches)
    create_card(slide4, Inches(0.6), Inches(1.75), Inches(6.0), Inches(5.0))
    add_card_header(slide4, Inches(0.8), Inches(1.9), Inches(5.6), "Radar Backscatter Mechanisms", color=C_CYAN, size=13)

    tb = slide4.shapes.add_textbox(Inches(0.8), Inches(2.35), Inches(5.6), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    sar_points = [
        ("• C-Band Active Microwave (5.405 GHz): ", "Transmits its own electromagnetic pulses, allowing day and night observation independent of solar illumination and penetrating rain/clouds."),
        ("• Specular Reflection on Water: ", "Smooth standing water acts as a specular reflector, bouncing radar pulses away from the sensor. Water thus appears distinctly dark (< -18 dB) in SAR imagery."),
        ("• Diffuse Scattering on Land: ", "Rough soil, urban structures, and terrain scatter energy back to the antenna in all directions, yielding bright backscatter values (-12 dB to -5 dB)."),
        ("• VV Polarization Channel: ", "Vertical transmit / Vertical receive. Highly sensitive to surface water roughness, capillary waves, and open water boundaries."),
        ("• VH Polarization Channel: ", "Vertical transmit / Horizontal receive (Cross-polarization). Captures depolarized volume scattering from vegetation, distinguishing crops from inundated fields.")
    ]
    for i, (bold_p, text_p) in enumerate(sar_points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = bold_p
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = text_p
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    # Right Column: Visual of real sample
    create_card(slide4, Inches(6.8), Inches(1.75), Inches(5.933), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(slide4, Inches(7.0), Inches(1.9), Inches(5.5), "Verified Dual-Polarization Dataset Sample", color=C_GREEN, size=13)

    sample_img = os.path.join(assets_dir, "outputs/visualizations/dataset_samples/Ghana_313799_verification.png")
    if os.path.exists(sample_img):
        slide4.shapes.add_picture(sample_img, Inches(7.0), Inches(2.3), width=Inches(5.5))
        tb_cap = slide4.shapes.add_textbox(Inches(7.0), Inches(6.25), Inches(5.5), Inches(0.4))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Sen1Floods11 Ground Truth Verification: Raw VV/VH Channels vs Discrete Mask"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(slide4, """
Let us discuss the physics of Synthetic Aperture Radar.
Sentinel-1 operates in the C-band microwave spectrum at 5.405 GHz. Because radar is an active sensor, it emits its own microwave signal and measures the returned backscatter.
When radar pulses hit smooth, standing water, they undergo specular reflection—the pulses bounce away from the satellite antenna, making water appear characteristically dark, typically below minus 18 decibels.
In contrast, rough land, urban structures, and dry soils scatter energy diffusely back to the sensor, appearing bright.
Furthermore, we use dual polarization:
The VV channel is vertically polarized and is sensitive to surface water roughness.
The VH cross-polarization channel captures depolarized volume scattering from vegetation, helping the model differentiate dry crops from flooded terrain.
""")

    # =========================================================================
    # SLIDE 5: DATASET SEN1FLOODS11
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, C_BG)
    add_header(slide5, "Data Engineering & Audit", "Benchmark Dataset: Sen1Floods11", "Global distribution, split verification, and discrete ground truth label encoding")
    add_footer(slide5, 5)

    # Left Side: Split Cards
    w_left = Inches(6.0)
    create_card(slide5, Inches(0.6), Inches(1.75), w_left, Inches(2.4))
    add_card_header(slide5, Inches(0.8), Inches(1.9), w_left - Inches(0.4), "Official Dataset Partitions (Hand-Labeled)", color=C_CYAN, size=13)

    splits = [
        ("Training Split:", "252 chips", "(1,008 GeoTIFF files)"),
        ("Validation Split:", "89 chips", "(356 GeoTIFF files)"),
        ("Official Test Split:", "90 chips", "(360 GeoTIFF files)"),
        ("Bolivia Holdout:", "15 chips", "(60 GeoTIFF files)")
    ]
    tb = slide5.shapes.add_textbox(Inches(0.8), Inches(2.3), w_left - Inches(0.4), Inches(1.7))
    tf = tb.text_frame
    for i, (s_name, count, detail) in enumerate(splits):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {s_name:<20} "
        r1.font.bold = True
        r1.font.color.rgb = C_CYAN_LIGHT
        r1.font.size = Pt(11)
        r2 = p.add_run()
        r2.text = f"{count:<10} "
        r2.font.bold = True
        r2.font.color.rgb = C_WHITE
        r2.font.size = Pt(11)
        r3 = p.add_run()
        r3.text = detail
        r3.font.color.rgb = C_MUTED
        r3.font.size = Pt(10)

    # Data Leakage & Label Card
    create_card(slide5, Inches(0.6), Inches(4.3), w_left, Inches(2.45))
    add_card_header(slide5, Inches(0.8), Inches(4.45), w_left - Inches(0.4), "Label Encoding & Data Leakage Audit", color=C_GREEN, size=13)

    tb_lbl = slide5.shapes.add_textbox(Inches(0.8), Inches(4.85), w_left - Inches(0.4), Inches(1.8))
    tf = tb_lbl.text_frame
    lbl_pts = [
        ("• Zero Data Leakage: ", "Verified pairwise mutual exclusivity across all CSV splits — 0 overlapping chips detected."),
        ("• Label -1 (Invalid/Ignored): ", "Missing data, cloud shadows, or sensor artifacts (masked out during loss computation)."),
        ("• Label 0 (Dry Land): ", "Uninundated terrain, buildings, soil, roads, and non-flooded vegetation."),
        ("• Label 1 (Water/Flood): ", "Inundated floodplains, standing water, and open water bodies (12.51% of valid test pixels)."),
        ("• Baseline Permanence: ", "JRC 30-year permanence raster provides historical water baseline (no separate pre-flood SAR files).")
    ]
    for i, (bp, tp) in enumerate(lbl_pts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = bp
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_GREEN if "Zero" in bp else (C_AMBER if "-1" in bp else C_WHITE)
        r2 = p.add_run()
        r2.text = tp
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_TEXT_LIGHT

    # Right Side: Verification sample image
    create_card(slide5, Inches(6.8), Inches(1.75), Inches(5.933), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(slide5, Inches(7.0), Inches(1.9), Inches(5.5), "Bolivia Event Sample Verification", color=C_WHITE, size=13)

    bol_img = os.path.join(assets_dir, "outputs/visualizations/dataset_samples/Bolivia_103757_verification.png")
    if os.path.exists(bol_img):
        slide5.shapes.add_picture(bol_img, Inches(7.0), Inches(2.3), width=Inches(5.5))
        tb_cap = slide5.shapes.add_textbox(Inches(7.0), Inches(6.25), Inches(5.5), Inches(0.4))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Bolivia Holdout Verification: S1 VV, S1 VH, S2 Optical, and Hand-Labeled Target"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(slide5, """
Here we present our benchmark dataset: Sen1Floods11, created by Cloud to Street and published at IEEE CVPRW.
It encompasses 11 catastrophic flood events across 6 continents—including the USA, Pakistan, Bolivia, Ghana, Mekong, Nigeria, Somalia, Spain, and Sri Lanka.
We strictly partitioned the dataset into:
252 training chips, 89 validation chips, 90 official test chips, and 15 Bolivia holdout chips.
Crucially, we performed a programmatic data leakage audit, proving zero overlapping chips between training and test sets.
The ground truth masks contain three discrete values:
Minus 1 for invalid or corrupted pixels, 0 for dry land, and 1 for floodwater.
We also clarify that Sen1Floods11 does not contain separate pre-flood SAR files; the historical surface water baseline is established via the JRC Global Surface Water 30-year permanence raster.
""")

    # =========================================================================
    # SLIDE 6: DATA PREPROCESSING
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, C_BG)
    add_header(slide6, "Pipeline Engineering", "Data Preprocessing & Tensor Pipeline", "End-to-end normalization, strided spatial subsampling, and synchronized augmentation")
    add_footer(slide6, 6)

    # 5 Horizontal Flow Step Cards
    step_w = Inches(2.25)
    step_h = Inches(4.8)
    gap = Inches(0.2)
    start_x = Inches(0.6)

    steps = [
        ("Step 1", "Decibel Clipping", C_CYAN, [
            "• Input: Raw Float32 GeoTIFFs.",
            "• Range: [-35 dB, +5 dB].",
            "• Extreme specular outliers below -35 dB and corner reflectors above +5 dB clipped.",
            "• Prevents radar flare distortions."
        ]),
        ("Step 2", "Min-Max Scaling", C_CYAN_LIGHT, [
            "• Target Range: [0.0, 1.0].",
            "• Formula: (dB - (-35)) / 40.0.",
            "• Normalizes both VV and VH channels identically.",
            "• Ensures smooth gradient flow during backprop."
        ]),
        ("Step 3", "Strided Subsampling", C_GREEN, [
            "• Resolution: 256 × 256.",
            "• Method: Slicing ([::stride, ::stride]).",
            "• Preserves discrete integer labels (-1, 0, 1).",
            "• Avoids bilinear/bicubic interpolation blur."
        ]),
        ("Step 4", "Synchronized Augment", C_AMBER, [
            "• Random Horizontal Flip (p=0.5).",
            "• Random Vertical Flip (p=0.5).",
            "• Random 90° Rotations (p=0.5).",
            "• Applied identically to SAR channels and mask."
        ]),
        ("Step 5", "Tensor Assembly", C_WHITE, [
            "• Tensor Shape: (2, 256, 256).",
            "• Channel 0: Normalized VV.",
            "• Channel 1: Normalized VH.",
            "• Target Tensor: (1, 256, 256) LongTensor.",
            "• Ready for U-Net forward pass."
        ])
    ]

    for i, (s_num, s_title, s_color, s_bullets) in enumerate(steps):
        x = start_x + i * (step_w + gap)
        create_card(slide6, x, Inches(1.85), step_w, step_h)

        # Step header badge
        badge_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), Inches(2.0), step_w - Inches(0.4), Inches(0.32))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = s_color
        badge_box.line.fill.background()
        tf = badge_box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = s_num.upper()
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_BG if s_color != C_CARD else C_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Step title
        tb = slide6.shapes.add_textbox(x + Inches(0.15), Inches(2.4), step_w - Inches(0.3), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

        # Bullets
        tb_b = slide6.shapes.add_textbox(x + Inches(0.15), Inches(2.9), step_w - Inches(0.3), Inches(3.6))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        for j, b_text in enumerate(s_bullets):
            p = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            p.text = b_text
            p.font.size = Pt(10)
            p.font.color.rgb = C_TEXT_LIGHT
            p.space_after = Pt(6)

    add_speaker_notes(slide6, """
Slide 6 illustrates our 5-stage data preprocessing pipeline.
Stage 1: Raw Sentinel-1 GeoTIFFs contain decibel values with extreme outliers. We clip the backscatter to the physical range of minus 35 to plus 5 decibels.
Stage 2: We linearly normalize this range to [0.0, 1.0], ensuring numerical stability for PyTorch tensors.
Stage 3: To downsample the satellite chips to 256 by 256 resolution, we intentionally avoid bilinear or bicubic interpolation on the masks, which would corrupt discrete integer labels. Instead, we use strided slicing, preserving the exact -1, 0, and 1 values.
Stage 4: We apply synchronized data augmentations—including random horizontal and vertical flips and 90-degree orthogonal rotations—applied identically to both SAR channels and the ground truth.
Stage 5: We package the normalized VV and VH channels into a two-channel input tensor ready for the U-Net.
""")

    # =========================================================================
    # SLIDE 7: U-NET ARCHITECTURE
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, C_BG)
    add_header(slide7, "Model Design", "Customized PyTorch U-Net Architecture", "Encoder-decoder segmentation network with Group Normalization (G=8) and skip connections")
    add_footer(slide7, 7)

    # Left: Architecture Flow Block (7.0 inches)
    create_card(slide7, Inches(0.6), Inches(1.75), Inches(7.2), Inches(5.0))
    add_card_header(slide7, Inches(0.8), Inches(1.9), Inches(6.8), "Hierarchical Feature Flow & Channel Dimensions", color=C_CYAN, size=13)

    # Visual blocks representing U-Net
    # Input
    in_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.35), Inches(6.6), Inches(0.4))
    in_box.fill.solid()
    in_box.fill.fore_color.rgb = C_CARD_DARK
    in_box.line.color.rgb = C_CYAN
    tf = in_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "INPUT TENSOR: 2 × 256 × 256 (Channel 0: S1-VV, Channel 1: S1-VH)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.alignment = PP_ALIGN.CENTER

    # Encoder stages
    enc_stages = [
        ("Encoder Stage 1: Conv3x3 (2 → 32) + GroupNorm(G=8) + GELU + MaxPool2d", "32 × 128 × 128"),
        ("Encoder Stage 2: Conv3x3 (32 → 64) + GroupNorm(G=8) + GELU + MaxPool2d", "64 × 64 × 64"),
        ("Encoder Stage 3: Conv3x3 (64 → 128) + GroupNorm(G=8) + GELU + MaxPool2d", "128 × 32 × 32"),
        ("Encoder Stage 4: Conv3x3 (128 → 256) + GroupNorm(G=8) + GELU + MaxPool2d", "256 × 16 × 16")
    ]
    for i, (txt, dim) in enumerate(enc_stages):
        box = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.85 + i * 0.42), Inches(5.1), Inches(0.35))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_DARK
        box.line.color.rgb = C_BORDER
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"↓ {txt}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_TEXT_LIGHT

        dim_box = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.05), Inches(2.85 + i * 0.42), Inches(1.45), Inches(0.35))
        dim_box.fill.solid()
        dim_box.fill.fore_color.rgb = C_BORDER_CYAN
        dim_box.line.fill.background()
        tf2 = dim_box.text_frame
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p2 = tf2.paragraphs[0]
        p2.text = dim
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE
        p2.alignment = PP_ALIGN.CENTER

    # Bottleneck
    b_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(4.55), Inches(6.6), Inches(0.42))
    b_box.fill.solid()
    b_box.fill.fore_color.rgb = RGBColor(124, 58, 237) # Purple accent
    b_box.line.fill.background()
    tf = b_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "★ BOTTLENECK: Conv3x3 (256 → 512) + Spatial Dropout (p=0.10) [512 × 16 × 16]"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Decoder summary
    dec_box = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.05), Inches(6.6), Inches(0.85))
    dec_box.fill.solid()
    dec_box.fill.fore_color.rgb = C_CARD_DARK
    dec_box.line.color.rgb = C_GREEN
    tf = dec_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "↑ DECODER STAGES: 4 Expanding Blocks [256 → 128 → 64 → 32]"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_GREEN
    p2 = tf.add_paragraph()
    p2.text = "• Bilinear Upsampling + Skip-Connection Concatenation + DoubleConv + GroupNorm"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_LIGHT

    # Output
    out_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.0), Inches(6.6), Inches(0.4))
    out_box.fill.solid()
    out_box.fill.fore_color.rgb = C_GREEN
    out_box.line.fill.background()
    tf = out_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "OUTPUT: 1×1 Conv (32 → 1) → Single-Channel Logits (1 × 256 × 256)"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_BG
    p.alignment = PP_ALIGN.CENTER

    # Right: Key Engineering Innovations (4.7 inches)
    create_card(slide7, Inches(8.0), Inches(1.75), Inches(4.733), Inches(5.0))
    add_card_header(slide7, Inches(8.2), Inches(1.9), Inches(4.3), "Key Architecture Innovations", color=C_GREEN, size=13)

    tb_spec = slide7.shapes.add_textbox(Inches(8.2), Inches(2.35), Inches(4.3), Inches(4.2))
    tf_spec = tb_spec.text_frame
    tf_spec.word_wrap = True
    specs = [
        ("GroupNorm (G=8) vs BatchNorm:", "BatchNorm computes statistics across the batch dimension. With small satellite batch sizes (B=2), BatchNorm exhibits extreme variance and instability. GroupNorm divides 32-512 channels into 8 independent groups per sample, providing rock-solid stability."),
        ("Bilinear Upsampling:", "Replaces transposed convolutions to eliminate checkerboard deconvolution artifacts in fine riverbank boundaries."),
        ("Skip Connections:", "Directly concatenates high-resolution spatial features from encoder to decoder, preserving crisp water boundaries."),
        ("Parameter Count: 7,760,257", "Balanced model capacity offering rich representational power while remaining lightweight enough for real-time edge/web inference (<50ms).")
    ]
    for i, (head, body) in enumerate(specs):
        p = tf_spec.paragraphs[0] if i == 0 else tf_spec.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {head}\n  "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = body
        r2.font.size = Pt(9.8)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(8)

    add_speaker_notes(slide7, """
Slide 7 details our core neural network architecture.
We implemented an optimized PyTorch U-Net tailored for 2-channel SAR inputs.
The encoder consists of 4 contracting stages starting at 32 feature channels and doubling to 256 channels at 16 by 16 spatial resolution.
The bottleneck expands to 512 channels and incorporates spatial dropout with rate 0.10 to prevent overfitting.
The decoder uses bilinear upsampling and skip connections to restore full 256 by 256 spatial detail.
A crucial viva point: why Group Normalization?
With satellite data, GPU memory constraints force small batch sizes (batch size = 2). Standard Batch Normalization calculates running statistics across the batch, leading to noisy, unstable gradients. Group Normalization with G=8 groups operates independently per image, guaranteeing batch-size invariant stability.
Total trainable parameters are exactly 7,760,257.
""")

    # =========================================================================
    # SLIDE 8: LOSS FUNCTION & TRAINING
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, C_BG)
    add_header(slide8, "Optimization & Convergence", "Loss Formulation & Training Configuration", "Hybrid masked loss overcoming class imbalance and verified convergence curves")
    add_footer(slide8, 8)

    # Left Card: Loss Function
    w_half = Inches(5.9)
    create_card(slide8, Inches(0.6), Inches(1.75), w_half, Inches(5.0))
    add_card_header(slide8, Inches(0.8), Inches(1.9), w_half - Inches(0.4), "Domain-Specific Masked Loss Formulation", color=C_CYAN, size=13)

    tb_loss = slide8.shapes.add_textbox(Inches(0.8), Inches(2.35), w_half - Inches(0.4), Inches(4.2))
    tf = tb_loss.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Loss Formula (Evaluated Strictly on Valid Pixels y ≠ -1):"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = "ℒ_total = 0.5 × ℒ_BCE + 0.5 × ℒ_Dice"
    p2.font.name = "Courier New"
    p2.font.bold = True
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_CYAN
    p2.space_before = Pt(4)
    p2.space_after = Pt(8)

    loss_details = [
        ("1. Binary Cross-Entropy (BCE):", "Calculates pixel-level logarithmic divergence between predicted logits and true labels, ensuring smooth probabilistic convergence."),
        ("2. Soft Dice Loss:", "Formulated as 1 - (2|X ∩ Y| + ε) / (|X| + |Y| + ε). Directly optimizes segmentation overlap, countering severe class imbalance where water occupies only 12.5% of pixels."),
        ("3. Strict Masking:", "Pixels labeled -1 (cloud shadows, sensor anomalies) are dynamically masked out in PyTorch tensors prior to loss reduction, preventing corrupted gradient backpropagation.")
    ]
    for h, b in loss_details:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{h} "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    # Right Card: Training Setup + Real Loss Curve Image
    create_card(slide8, Inches(6.8), Inches(1.75), w_half + Inches(0.033), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(slide8, Inches(7.0), Inches(1.9), w_half - Inches(0.4), "Training Configuration & Real Loss Curve", color=C_GREEN, size=13)

    # Training specs summary pill
    spec_tb = slide8.shapes.add_textbox(Inches(7.0), Inches(2.25), w_half - Inches(0.4), Inches(0.9))
    tf_s = spec_tb.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "• Epochs: 5 (Early Stopping: 2)  |  Batch Size: 2  |  Resolution: 256×256\n• Optimizer: AdamW (lr = 0.0005, weight_decay = 1e-4)  |  Cosine Annealing\n• Hardware: Apple Silicon MPS  |  Total Wall-Clock Time: ~15.2 seconds"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_TEXT_LIGHT

    loss_img = os.path.join(assets_dir, "outputs/training/loss_curves/s1_loss_curve.png")
    if os.path.exists(loss_img):
        slide8.shapes.add_picture(loss_img, Inches(7.0), Inches(3.2), width=Inches(5.5))
        tb_cap = slide8.shapes.add_textbox(Inches(7.0), Inches(6.3), Inches(5.5), Inches(0.3))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Actual Training & Validation Loss Curves (Smooth Cosine Decay)"
        p.font.size = Pt(9)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(slide8, """
Slide 8 covers our loss function formulation and optimization.
Because water represents only 12.5% of valid pixels across the dataset, standard cross-entropy loss would bias the network toward predicting dry land everywhere.
To overcome this severe class imbalance, we formulated a hybrid loss:
0.5 BCE plus 0.5 Soft Dice loss.
BCE provides smooth pixel-level classification learning, while Dice loss directly maximizes spatial region overlap.
Furthermore, invalid minus 1 pixels are dynamically excluded via boolean indexing before computing loss gradients.
For optimization, we used AdamW with an initial learning rate of 0.0005, weight decay of 1e-4, and Cosine Annealing learning rate scheduling.
The model converged in 5 epochs on Apple Silicon MPS in approximately 15.2 seconds total wall-clock time.
""")

    # =========================================================================
    # SLIDE 9: EVALUATION METRICS
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, C_BG)
    add_header(slide9, "Quantitative Assessment", "Segmentation Evaluation Metrics & Formulations", "Mathematical formulations evaluated strictly over valid pixels V = {i | y_i != -1}")
    add_footer(slide9, 9)

    # 4 Metric Cards in 2x2 Grid + Bottom Insight Banner
    c_w = Inches(5.8)
    c_h = Inches(1.9)

    metrics_list = [
        ("Intersection over Union (IoU / Jaccard Index)", "IoU = TP / (TP + FP + FN)", "Measures the geometric overlap between the predicted flood mask and ground truth mask. The gold standard for remote sensing segmentation.", C_CYAN, Inches(0.6), Inches(1.75)),
        ("Dice Coefficient / F1 Score", "Dice = 2·TP / (2·TP + FP + FN) = F1", "Harmonic mean of precision and recall. Directly penalizes both false alarms (FP) and missed floodwaters (FN).", C_GREEN, Inches(6.9), Inches(1.75)),
        ("Precision (Positive Predictive Value)", "Precision = TP / (TP + FP)", "What percentage of predicted flood pixels are genuine water? Critical for preventing false evacuation alarms.", C_CYAN_LIGHT, Inches(0.6), Inches(3.8)),
        ("Recall (Sensitivity / Detection Rate)", "Recall = TP / (TP + FN)", "What percentage of actual floodwaters were successfully detected? Essential for disaster rescue operations.", C_AMBER, Inches(6.9), Inches(3.8))
    ]

    for title, formula, desc, color, x, y in metrics_list:
        create_card(slide9, x, y, c_w, c_h)
        add_card_header(slide9, x + Inches(0.2), y + Inches(0.12), c_w - Inches(0.4), title, color=color, size=11.5)

        tb = slide9.shapes.add_textbox(x + Inches(0.2), y + Inches(0.45), c_w - Inches(0.4), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = formula
        p.font.name = "Courier New"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT

    # Bottom Banner: Why Accuracy is Deceptive
    create_card(slide9, Inches(0.6), Inches(5.85), Inches(12.1), Inches(0.95), bg_color=C_CARD_DARK, border_color=C_BORDER_CYAN)
    tb_bot = slide9.shapes.add_textbox(Inches(0.8), Inches(5.92), Inches(11.7), Inches(0.8))
    tf_b = tb_bot.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "CRITICAL VIVA INSIGHT — Why Overall Accuracy is Deceptive in Flood Mapping:"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_AMBER
    p2 = tf_b.add_paragraph()
    p2.text = "Because 87.5% of pixels are dry land, a naive model predicting 100% land would achieve 87.5% accuracy while completely failing to detect any floodwaters. Therefore, IoU (0.5548) and Dice (0.7137) are the true benchmarks of segmentation efficacy."
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(slide9, """
Slide 9 defines our quantitative evaluation formulations.
All metrics are computed exclusively over valid pixels, strictly ignoring -1 invalid labels:
IoU, or the Jaccard Index: True Positives divided by True Positives plus False Positives plus False Negatives. This is the primary standard in computer vision benchmarks.
Dice Score / F1: Two times True Positives over total predicted and actual positives.
Precision: True Positives over all predicted positives—critical for minimizing false alarms.
Recall: True Positives over all actual flood pixels—essential for not missing inundated communities.
A classic viva question is: 'Why not simply use Overall Accuracy?'
As shown on the slide, dry land accounts for 87.5% of the pixels. A trivial model predicting only dry land would achieve 87.5% accuracy with zero flood detection capability. Hence, IoU and Dice are the mathematically rigorous metrics.
""")

    # =========================================================================
    # SLIDE 10: OFFICIAL TEST RESULTS
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, C_BG)
    add_header(slide10, "Quantitative Benchmark", "Official Complete Test Results (90 / 90 Chips)", "Verified performance evaluated across 100% of the official Sen1Floods11 test split")
    add_footer(slide10, 10)

    # 5 Large KPI Cards across top
    kpis = [
        ("TEST IoU", "0.5548", "Jaccard Index", C_CYAN),
        ("DICE / F1", "0.7137", "Overlap Coeff", C_GREEN),
        ("PRECISION", "0.7717", "77.17% Water True", C_CYAN_LIGHT),
        ("RECALL", "0.6637", "66.37% Water Found", C_AMBER),
        ("ACCURACY", "0.9334", "93.34% Pixel Acc", C_WHITE)
    ]
    card_w = Inches(2.25)
    for i, (label, val, sub, col) in enumerate(kpis):
        x = Inches(0.6 + i * 2.47)
        create_card(slide10, x, Inches(1.75), card_w, Inches(1.5), bg_color=C_CARD, border_color=C_BORDER)

        tb = slide10.shapes.add_textbox(x + Inches(0.1), Inches(1.85), card_w - Inches(0.2), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = label
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.alignment = PP_ALIGN.CENTER

        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.alignment = PP_ALIGN.CENTER

    # Bottom Left: Scope Summary Table
    create_card(slide10, Inches(0.6), Inches(3.45), Inches(5.9), Inches(3.3))
    add_card_header(slide10, Inches(0.8), Inches(3.6), Inches(5.5), "Full Benchmark Evaluation Scope", color=C_CYAN, size=13)

    tb_scope = slide10.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(5.5), Inches(2.6))
    tf = tb_scope.text_frame
    scope_rows = [
        ("• Test Chips Evaluated:", "90 / 90 (100% complete)"),
        ("• Total Valid Pixels:", "20,517,367 pixels"),
        ("• Ground Truth Water Pixels:", "2,566,101 pixels (12.51%)"),
        ("• Ground Truth Land Pixels:", "17,951,266 pixels (87.49%)"),
        ("• Test Masked Loss:", "0.3918"),
        ("• Inference Speed:", "< 50ms per 256×256 chip")
    ]
    for i, (k, v) in enumerate(scope_rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{k:<28} "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = v
        r2.font.bold = True
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = C_WHITE

    # Bottom Right: Real Metric Curve Image
    create_card(slide10, Inches(6.8), Inches(3.45), Inches(5.933), Inches(3.3), bg_color=C_CARD_DARK)
    add_card_header(slide10, Inches(7.0), Inches(3.6), Inches(5.5), "Verified Validation Metric Progression", color=C_GREEN, size=13)

    metric_img = os.path.join(assets_dir, "outputs/training/metric_curves/s1_metric_curve.png")
    if os.path.exists(metric_img):
        slide10.shapes.add_picture(metric_img, Inches(7.0), Inches(4.0), width=Inches(5.5))

    add_speaker_notes(slide10, """
Slide 10 presents our official, complete test results.
Unlike partial evaluations, we evaluated all 90 chips in the official Sen1Floods11 test split without excluding difficult cases or cloud-affected scenes.
This encompasses over 20.5 million valid pixels.
Our verified results:
Test IoU: 0.5548
Test Dice / F1 Score: 0.7137
Test Precision: 0.7717
Test Recall: 0.6637
Overall Accuracy: 0.9334
Test Loss: 0.3918.
The high precision of 77.17% confirms that when our model identifies an area as flooded, it is highly dependable.
On the right, you can see our actual validation curves demonstrating rapid, stable convergence.
""")

    # =========================================================================
    # SLIDE 11: CONFUSION MATRIX
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, C_BG)
    add_header(slide11, "Error Diagnostics", "Pixel-Level Test Confusion Matrix", "Complete 20,517,367 pixel breakdown and error taxonomy across 90 test chips")
    add_footer(slide11, 11)

    # Left: 2x2 Matrix Graphic (6.0 inches)
    create_card(slide11, Inches(0.6), Inches(1.75), Inches(6.0), Inches(5.0))
    add_card_header(slide11, Inches(0.8), Inches(1.9), Inches(5.6), "Official Test Set Confusion Matrix", color=C_CYAN, size=13)

    # Matrix Table / Cards
    # TP
    tp_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.4), Inches(2.6), Inches(1.3))
    tp_box.fill.solid()
    tp_box.fill.fore_color.rgb = RGBColor(6, 78, 59) # Deep green
    tp_box.line.color.rgb = C_GREEN
    tf = tp_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRUE POSITIVES (TP)\n1,703,186 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Water correctly detected (66.4%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = C_GREEN
    p2.alignment = PP_ALIGN.CENTER

    # FP
    fp_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.7), Inches(2.4), Inches(2.6), Inches(1.3))
    fp_box.fill.solid()
    fp_box.fill.fore_color.rgb = RGBColor(127, 29, 29) # Deep red
    fp_box.line.color.rgb = RGBColor(239, 68, 68)
    tf = fp_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FALSE POSITIVES (FP)\n503,818 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Dry land flagged as flood (2.8%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = RGBColor(252, 165, 165)
    p2.alignment = PP_ALIGN.CENTER

    # FN
    fn_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.9), Inches(2.6), Inches(1.3))
    fn_box.fill.solid()
    fn_box.fill.fore_color.rgb = RGBColor(120, 53, 15) # Deep amber
    fn_box.line.color.rgb = C_AMBER
    tf = fn_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FALSE NEGATIVES (FN)\n862,915 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Floodwater missed (33.6%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = C_AMBER
    p2.alignment = PP_ALIGN.CENTER

    # TN
    tn_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.7), Inches(3.9), Inches(2.6), Inches(1.3))
    tn_box.fill.solid()
    tn_box.fill.fore_color.rgb = C_CARD_DARK
    tn_box.line.color.rgb = C_BORDER
    tf = tn_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRUE NEGATIVES (TN)\n17,447,448 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Land correctly identified (97.2%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = C_MUTED
    p2.alignment = PP_ALIGN.CENTER

    # Matrix summary note
    tb_ms = slide11.shapes.add_textbox(Inches(0.9), Inches(5.35), Inches(5.4), Inches(1.2))
    tf = tb_ms.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "• Total Evaluated Valid Pixels: 20,517,367 pixels (100% of 90 chips)\n• Ground Truth Water: 2,566,101 pixels (12.51%)\n• Ground Truth Land: 17,951,266 pixels (87.49%)"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_TEXT_LIGHT

    # Right: Error Diagnostics (5.9 inches)
    create_card(slide11, Inches(6.8), Inches(1.75), Inches(5.933), Inches(5.0))
    add_card_header(slide11, Inches(7.0), Inches(1.9), Inches(5.5), "Physical Radar Error Taxonomy", color=C_AMBER, size=13)

    tb_diag = slide11.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(5.5), Inches(4.2))
    tf = tb_diag.text_frame
    tf.word_wrap = True
    diag_pts = [
        ("1. Causes of False Positives (FP = 503,818):", [
            "• Smooth, flat dry terrain (e.g. airport runways, dry salt flats, calm sand).",
            "• Mirror-like terrain produces specular radar reflection resembling open water.",
            "• Mitigated by VH cross-polarization channel which captures texture."
        ]),
        ("2. Causes of False Negatives (FN = 862,915):", [
            "• Flooded vegetation and dense tree canopies causing corner-reflector 'double-bounce' backscatter.",
            "• Narrow drainage channels narrower than Sentinel-1's spatial resolution.",
            "• Wind-induced capillary waves roughening the water surface."
        ]),
        ("3. Disaster Management Impact:", [
            "• 77.17% Precision prevents false alarms and wasted relief resources.",
            "• 66.37% Recall captures major continuous floodplains reliably."
        ])
    ]
    for section_title, bullets in diag_pts:
        p = tf.add_paragraph()
        p.text = section_title
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_CYAN_LIGHT
        p.space_before = Pt(4)
        for b in bullets:
            p2 = tf.add_paragraph()
            p2.text = b
            p2.font.size = Pt(9.5)
            p2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(slide11, """
Slide 11 presents our complete pixel-level confusion matrix across all 20.5 million test pixels.
We have:
1,703,186 True Positives,
503,818 False Positives,
862,915 False Negatives, and
17,447,448 True Negatives.
Analyzing the physics behind these numbers is vital for any viva examination:
Why do False Positives occur?
Extremely smooth dry land, airport tarmac, and calm dry playas specularly reflect radar pulses away, mimicking open water.
Why do False Negatives occur?
When floodwaters submerge dense forest or vegetation, the tree trunks and water surface form a 90-degree corner reflector. This causes 'double-bounce' backscatter, making flooded forest appear bright rather than dark.
Understanding these radar physics phenomena allows us to contextualize model behavior accurately.
""")

    # =========================================================================
    # SLIDE 12: PREDICTION RESULTS & CASE STUDIES
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, C_BG)
    add_header(slide12, "Qualitative Case Studies", "Representative Test Case Qualitative Analysis", "Visualizing Good (USA), Average (Pakistan), and Difficult (Sri Lanka) benchmark predictions")
    add_footer(slide12, 12)

    # 3 Column Case Study Cards
    col_w = Inches(3.8)
    card_h = Inches(5.0)

    cases = [
        ("Case 1: USA_905409", "Good Prediction", C_GREEN, "outputs/visualizations/final_test/1_good_prediction_USA_905409.png", [
            ("IoU: ", "0.9350", " | Dice: ", "0.9664"),
            ("Precision: ", "0.9720", " | Recall: ", "0.9609"),
            ("Characteristics:", " Open agricultural floodplain with strong specular reflection and high contrast.")
        ]),
        ("Case 2: Pakistan_694942", "Average Prediction", C_CYAN, "outputs/visualizations/final_test/2_average_prediction_Pakistan_694942.png", [
            ("IoU: ", "0.3829", " | Dice: ", "0.5538"),
            ("Precision: ", "0.4152", " | Recall: ", "0.8313"),
            ("Characteristics:", " Complex agrarian field boundaries with saturated mud-water transition zones.")
        ]),
        ("Case 3: Sri-Lanka_450918", "Difficult Prediction", C_AMBER, "outputs/visualizations/final_test/3_difficult_prediction_Sri-Lanka_450918.png", [
            ("IoU: ", "0.0086", " | Dice: ", "0.0171"),
            ("Precision: ", "0.2317", " | Recall: ", "0.0089"),
            ("Characteristics:", " Flooded vegetation canopy with radar double-bounce backscatter.")
        ])
    ]

    for i, (c_title, c_badge, c_color, c_img_rel, c_stats) in enumerate(cases):
        x = Inches(0.6 + i * 4.15)
        create_card(slide12, x, Inches(1.75), col_w, card_h)

        # Header Badge
        add_card_header(slide12, x + Inches(0.15), Inches(1.85), col_w - Inches(0.3), c_title, color=c_color, size=11.5)

        # Image
        c_img_path = os.path.join(assets_dir, c_img_rel)
        if os.path.exists(c_img_path):
            slide12.shapes.add_picture(c_img_path, x + Inches(0.15), Inches(2.25), width=col_w - Inches(0.3))

        # Stats Box
        tb = slide12.shapes.add_textbox(x + Inches(0.15), Inches(5.0), col_w - Inches(0.3), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True

        # Line 1: Badge
        p = tf.paragraphs[0]
        p.text = f"Category: {c_badge}"
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = c_color

        # Line 2 & 3: Stats
        for row in c_stats:
            p = tf.add_paragraph()
            if len(row) == 4:
                r1 = p.add_run()
                r1.text = row[0]
                r1.font.bold = True
                r1.font.size = Pt(9)
                r1.font.color.rgb = C_CYAN_LIGHT
                r2 = p.add_run()
                r2.text = row[1]
                r2.font.bold = True
                r2.font.size = Pt(9)
                r2.font.color.rgb = C_WHITE
                r3 = p.add_run()
                r3.text = row[2]
                r3.font.bold = True
                r3.font.size = Pt(9)
                r3.font.color.rgb = C_CYAN_LIGHT
                r4 = p.add_run()
                r4.text = row[3]
                r4.font.bold = True
                r4.font.size = Pt(9)
                r4.font.color.rgb = C_WHITE
            else:
                r1 = p.add_run()
                r1.text = row[0]
                r1.font.bold = True
                r1.font.size = Pt(8.5)
                r1.font.color.rgb = C_MUTED
                r2 = p.add_run()
                r2.text = row[1]
                r2.font.size = Pt(8.5)
                r2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(slide12, """
Slide 12 demonstrates representative qualitative test cases from the official benchmark.
We deliberately present a transparent distribution of performance:
Case 1: USA_905409 (Good Prediction) achieves an IoU of 0.9350 and Dice of 0.9664. Here, open agricultural floodplains provide crisp specular contrast, allowing our U-Net to segment water bodies with near-perfect alignment.
Case 2: Pakistan_694942 (Average Prediction) achieves an IoU of 0.3829 and Recall of 0.8313. The model successfully captures the core inundation but experiences edge uncertainty along saturated agricultural boundaries.
Case 3: Sri-Lanka_450918 (Difficult Prediction) has an IoU of 0.0086. We explicitly highlight this case because it illustrates the physical limitation of C-band SAR: dense tropical forest canopies create double-bounce scattering that obscures standing water beneath.
Presenting all three cases proves our model was tested thoroughly against real-world complexities.
""")

    # =========================================================================
    # SLIDE 13: FLOOD MAPPING & NOMINAL AREA
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, C_BG)
    add_header(slide13, "Geospatial Analysis", "Flood Mapping & Nominal Flooded Area Estimation", "Transforming continuous probability heatmaps into actionable disaster management metrics")
    add_footer(slide13, 13)

    # Left Side: Mapping Workflow & Scientific Terminology (6.0 inches)
    w_left = Inches(6.0)
    create_card(slide13, Inches(0.6), Inches(1.75), w_left, Inches(5.0))
    add_card_header(slide13, Inches(0.8), Inches(1.9), w_left - Inches(0.4), "From Continuous Logits to Flooded Extent", color=C_CYAN, size=13)

    tb_map = slide13.shapes.add_textbox(Inches(0.8), Inches(2.35), w_left - Inches(0.4), Inches(4.2))
    tf = tb_map.text_frame
    tf.word_wrap = True

    mapping_steps = [
        ("1. Flood Probability Map (p̂ ∈ [0, 1]):", "Sigmoid activation transforms raw logits into continuous water probability heatmaps. Note: termed 'Flood Probability Map' rather than confidence map."),
        ("2. Probability Thresholding (τ = 0.50):", "Pixels with p̂ ≥ τ are classified as flooded. Threshold is dynamically adjustable in our Streamlit UI (0.10 to 0.90)."),
        ("3. Spatial Flood Mask Generation:", "Binary segmentation mask isolating inundated water polygons."),
        ("4. Nominal Flooded Area Estimation (km²):", "Calculated as: Area ≈ N_flood × 100 m² / 10⁶ = km²."),
        ("5. Crucial Geospatial Caveat (Nominal Area):", "Sen1Floods11 rasters are in WGS84 (EPSG:4326) with angular resolution Δ ≈ 8.98×10⁻⁵°. Physical pixel width varies with cos(latitude). Area is thus documented as a 'Nominal Flooded Area' equatorial approximation (10m nominal resolution).")
    ]
    for h, b in mapping_steps:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{h}\n"
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_CYAN_LIGHT if "Caveat" not in h else C_AMBER
        r2 = p.add_run()
        r2.text = f"{b}\n"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT

    # Right Side: Regional Prediction Visual
    create_card(slide13, Inches(6.8), Inches(1.75), Inches(5.933), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(slide13, Inches(7.0), Inches(1.9), Inches(5.5), "Somalia Regional Flood Inundation Case", color=C_GREEN, size=13)

    som_img = os.path.join(assets_dir, "outputs/visualizations/predictions/Somalia_989553_regional_prediction.png")
    if os.path.exists(som_img):
        slide13.shapes.add_picture(som_img, Inches(7.0), Inches(2.3), width=Inches(5.5))
        tb_cap = slide13.shapes.add_textbox(Inches(7.0), Inches(6.25), Inches(5.5), Inches(0.4))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Regional Mapping Output: Input SAR, Ground Truth, Prediction & Model Overlay"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(slide13, """
Slide 13 details how we convert model predictions into geospatial flood mapping products.
Our U-Net outputs continuous probability values between 0 and 1, which we term the 'Flood Probability Map'.
By applying an adjustable decision threshold (defaulting to 0.50), we extract the binary flooded mask.
To provide emergency services with quantitative insights, we compute the 'Nominal Flooded Area' in square kilometers.
An important academic point regarding geospatial projections:
The Sen1Floods11 rasters are distributed in unprojected WGS84 coordinates (EPSG:4326) with angular pixel spacing. Because longitudinal meter width scales with the cosine of latitude, true pixel area varies by latitude. Therefore, we explicitly report 'Nominal Flooded Area' based on standard 10-meter nominal Sentinel-1 resolution (100 square meters per pixel).
""")

    # =========================================================================
    # SLIDE 14: STREAMLIT APPLICATION
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, C_BG)
    add_header(slide14, "Interactive Demonstration", "Interactive Streamlit Web Dashboard (`app.py`)", "Dual-mode web interface supporting benchmark exploration and custom GeoTIFF upload")
    add_footer(slide14, 14)

    # Left: App Architecture Cards (6.0 inches)
    w_left = Inches(6.0)
    create_card(slide14, Inches(0.6), Inches(1.75), w_left, Inches(5.0))
    add_card_header(slide14, Inches(0.8), Inches(1.9), w_left - Inches(0.4), "Application Architecture & Features", color=C_CYAN, size=13)

    tb_app = slide14.shapes.add_textbox(Inches(0.8), Inches(2.35), w_left - Inches(0.4), Inches(4.2))
    tf = tb_app.text_frame
    tf.word_wrap = True

    app_features = [
        ("• Mode 1 — Benchmark Test Explorer:", "Allows instant dropdown selection of all 90 official test chips. Renders side-by-side 5-panel diagnostics: SAR False-Color, Ground Truth, Predicted Flood, Overlay, and Probability Heatmap."),
        ("• Mode 2 — Custom GeoTIFF Upload:", "Enables operators to upload raw dual-polarization Sentinel-1 .tif files. Validates channel counts (≥ 2 bands), normalizes dB, executes inference, and exports flood maps."),
        ("• Real-Time Threshold Slider:", "Dynamic threshold control (0.10 to 0.90) updating flood area and confusion metrics instantly in the browser."),
        ("• Sub-50ms Inference Latency:", "PyTorch model weights cached via @st.cache_resource, enabling immediate sub-50ms inference per chip."),
        ("• Live Demo Link:", "Operational at http://localhost:8501 for immediate examiner interaction.")
    ]
    for h, b in app_features:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{h}\n  "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT if "Demo" not in h else C_GREEN
        r2 = p.add_run()
        r2.text = f"{b}\n"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT

    # Right: Streamlit Demo Visual Frame
    create_card(slide14, Inches(6.8), Inches(1.75), Inches(5.933), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(slide14, Inches(7.0), Inches(1.9), Inches(5.5), "Live Dashboard Interface & Diagnostic Panels", color=C_GREEN, size=13)

    pred_demo = os.path.join(assets_dir, "outputs/visualizations/predictions/USA_994009_regional_prediction.png")
    if os.path.exists(pred_demo):
        slide14.shapes.add_picture(pred_demo, Inches(7.0), Inches(2.3), width=Inches(5.5))

        # Live Demo Callout Box
        demo_btn = slide14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(5.85), Inches(4.5), Inches(0.6))
        demo_btn.fill.solid()
        demo_btn.fill.fore_color.rgb = C_BORDER_CYAN
        demo_btn.line.fill.background()
        tf = demo_btn.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "LIVE SYSTEM DEMO → http://localhost:8501"
        p.font.name = FONT_HEADING
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(slide14, """
Slide 14 introduces our interactive demonstration application, developed with Streamlit in app.py.
The application operates in two distinct modes:
Mode 1 is the Benchmark Test Explorer, allowing users and examiners to select any of the 90 official test chips from a dropdown menu. The UI instantly visualizes the input SAR composite, ground truth mask, predicted flood mask, color-coded overlay, and continuous probability heatmap, alongside quantitative IoU, Dice, Precision, and Recall scores.
Mode 2 allows uploading arbitrary Sentinel-1 GeoTIFF rasters for zero-shot inference.
Inference takes under 50 milliseconds per chip thanks to model caching via st.cache_resource.
We are prepared to transition to the live demo at localhost:8501 during questions.
""")

    # =========================================================================
    # SLIDE 15: LIMITATIONS & FUTURE SCOPE
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, C_BG)
    add_header(slide15, "Critical Analysis & Roadmap", "Academic Limitations & Future Scope", "Objective evaluation of current constraints and proposed engineering enhancements")
    add_footer(slide15, 15)

    # 2 Column Cards
    col_w = Inches(5.9)
    card_h = Inches(5.0)

    # Left: Limitations
    create_card(slide15, Inches(0.6), Inches(1.75), col_w, card_h)
    add_card_header(slide15, Inches(0.8), Inches(1.9), col_w - Inches(0.4), "Current Academic Limitations", color=C_AMBER, size=13)

    tb_lim = slide15.shapes.add_textbox(Inches(0.8), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb_lim.text_frame
    tf.word_wrap = True
    limits = [
        ("• Specular Surface False Positives: ", "Smooth dry asphalt, airport runways, and flat desert sand specularly reflect radar, causing occasional false alarms."),
        ("• Canopy Volume False Negatives: ", "Flooded vegetation and dense forests induce double-bounce backscatter, causing radar energy to return bright and obscuring standing water."),
        ("• Unlabeled No-Data Pixels: ", "Sen1Floods11 contains -1 invalid pixels which must be carefully masked out during training."),
        ("• Historical JRC Baseline: ", "Sen1Floods11 relies on the 30-year JRC permanence raster rather than paired pre-flood SAR acquisitions."),
        ("• Nominal Equatorial Area Approximation: ", "WGS84 angular resolution requires nominal metric assumptions rather than local UTM projection.")
    ]
    for h, b in limits:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_AMBER
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(4)

    # Right: Future Scope
    create_card(slide15, Inches(6.8), Inches(1.75), col_w + Inches(0.033), card_h, border_color=C_BORDER_CYAN)
    add_card_header(slide15, Inches(7.0), Inches(1.9), col_w - Inches(0.4), "Future Research & Engineering Scope", color=C_CYAN, size=13)

    tb_fut = slide15.shapes.add_textbox(Inches(7.0), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb_fut.text_frame
    tf.word_wrap = True
    futures = [
        ("• Bi-Temporal SAR Difference Networks: ", "Ingesting co-registered pre-flood and post-flood SAR pairs to directly isolate floodwater from permanent lakes without optical baselines."),
        ("• Multi-Sensor & DEM Fusion: ", "Fusing Sentinel-1 SAR with Copernicus DEM elevation rasters and Sentinel-2 optical data to filter low-lying flood-prone areas."),
        ("• Critical Infrastructure Damage Intersection: ", "Overlaying OpenStreetMap building footprints, hospitals, and road networks onto predicted flood masks for automated evacuation routing."),
        ("• Cloud-Native Automated Alerting: ", "Deploying as a serverless container hooked to Copernicus Open Access Hub webhooks for automated near-real-time flood alerts.")
    ]
    for h, b in futures:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    add_speaker_notes(slide15, """
Slide 15 provides an objective assessment of our system's limitations and future roadmap.
Limitations:
1. False positives on smooth dry land like airport runways due to specular radar reflection.
2. False negatives in dense forests where double-bounce backscatter masks standing water.
3. The benchmark uses the 30-year JRC permanence raster as baseline rather than paired pre-flood SAR acquisitions.
4. Area calculation is a nominal equatorial approximation due to unprojected WGS84 rasters.
Future Scope:
1. Developing bi-temporal SAR difference networks using paired pre-and-post flood SAR passes.
2. Fusing DEM topography data to enforce hydrological gravity constraints.
3. Intersecting predicted flood polygons with OpenStreetMap road and building vector layers to calculate stranded populations and blocked evacuation routes.
""")

    # =========================================================================
    # SLIDE 16: CONCLUSION
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide16, C_BG)

    # Outer decorative card
    create_card(slide16, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3), bg_color=C_CARD_DARK, border_color=C_GREEN, line_width=1.5)

    # Title Badge
    badge = slide16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.9), Inches(3.2), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_GREEN
    badge.line.fill.background()
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "CONCLUSION & VIVA DEFENSE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_BG
    p.alignment = PP_ALIGN.CENTER

    # Main Heading
    tb_title = slide16.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(11.0), Inches(0.6))
    tf = tb_title.text_frame
    p = tf.paragraphs[0]
    p.text = "From Satellite Radar to Verified Flood Maps"
    p.font.name = FONT_HEADING
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # 4 Key Achievement Summary Boxes
    box_w = Inches(2.6)
    box_h = Inches(2.2)
    achievements = [
        ("ROBUST PIPELINE", "Sentinel-1 VV/VH", "Automated decibel normalization & strided subsampling pipeline with 0 data leakage.", C_CYAN),
        ("OPTIMIZED U-NET", "7.76M Parameters", "Group Normalization (G=8) ensured stable convergence with batch size 2.", C_GREEN),
        ("TEST BENCHMARK", "0.5548 Test IoU", "Evaluated on 100% of 90 test chips (20.5M pixels) with 0.7137 Dice & 93.34% Acc.", C_CYAN_LIGHT),
        ("VERIFIED QUALITY", "38/38 Tests Passed", "Full test suite passing, live Streamlit UI & sub-50ms inference latency.", C_WHITE)
    ]
    for i, (head, stat, desc, col) in enumerate(achievements):
        x = Inches(1.0 + i * 2.8)
        create_card(slide16, x, Inches(2.1), box_w, box_h, bg_color=C_CARD, border_color=C_BORDER)

        tb = slide16.shapes.add_textbox(x + Inches(0.1), Inches(2.2), box_w - Inches(0.2), box_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = stat
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(9)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.alignment = PP_ALIGN.CENTER

    # Thank You & Q&A Box
    tb_ty = slide16.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.333), Inches(1.8))
    tf = tb_ty.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Thank You for Your Time & Consideration"
    p.font.name = FONT_HEADING
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Open for Questions, Evaluation, and Live Demonstration"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)

    p3 = tf.add_paragraph()
    p3.text = "Project Repository: AI-Based Flood Detection and Mapping  |  Sen1Floods11 • Sentinel-1 SAR • PyTorch U-Net"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(10)
    p3.font.color.rgb = C_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(10)

    add_speaker_notes(slide16, """
In conclusion, we have successfully developed, audited, and verified an end-to-end deep learning semantic segmentation system for satellite SAR flood mapping.
By leveraging Sentinel-1 C-band dual polarization, Group Normalization with 8 groups, and a masked hybrid BCE-Dice loss, we achieved a Test IoU of 0.5548 and Dice score of 0.7137 across all 90 chips of the official benchmark without data leakage.
Our complete test suite of 38 unit tests passes with 100% success, and our Streamlit application provides real-time flood mapping and GeoTIFF analysis.
Thank you for your time and guidance. We now welcome questions and are ready to proceed with the live demonstration.
""")

    # =========================================================================
    # SAVE PRESENTATION
    # =========================================================================
    output_path = os.path.join(assets_dir, "AI_Flood_Detection_Minor_Project.pptx")
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides!")
    return output_path

if __name__ == "__main__":
    build_presentation()
