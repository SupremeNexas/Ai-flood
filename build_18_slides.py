"""
build_18_slides.py
Generates the revised, highly legible 18-slide presentation:
AI_Flood_Detection_Minor_Project_Final.pptx
and updates AI_Flood_Detection_Minor_Project.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Color Palette ---
C_BG = RGBColor(11, 19, 43)             # #0B132B Deep Navy Background
C_CARD = RGBColor(30, 41, 59)           # #1E293B Slate Card Fill
C_CARD_DARK = RGBColor(15, 23, 42)      # #0F172A Darker Slate Card
C_BORDER = RGBColor(51, 65, 85)         # #334155 Slate Border
C_BORDER_CYAN = RGBColor(2, 132, 199)   # #0284C7 Accent Cyan Border
C_CYAN = RGBColor(0, 229, 255)          # #00E5FF Bright Electric Cyan
C_CYAN_LIGHT = RGBColor(56, 189, 248)   # #38BDF8 Sky Cyan Accent
C_GREEN = RGBColor(16, 185, 129)        # #10B981 Emerald Green
C_GREEN_DARK = RGBColor(6, 78, 59)      # #064E3B Deep Emerald
C_AMBER = RGBColor(245, 158, 11)        # #F59E0B Amber / Warning
C_RED = RGBColor(239, 68, 68)           # #EF4444 Red Accent
C_RED_DARK = RGBColor(127, 29, 29)      # #7F1D1D Deep Red
C_PURPLE = RGBColor(139, 92, 246)       # #8B5CF6 Purple Accent
C_WHITE = RGBColor(255, 255, 255)       # #FFFFFF Pure White
C_MUTED = RGBColor(148, 163, 184)       # #94A3B8 Muted Slate Gray
C_TEXT_LIGHT = RGBColor(241, 245, 249)  # #F1F5F9 Off-white Body

FONT_HEADING = "Arial"
FONT_BODY = "Arial"

def set_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = C_BG

def add_header(slide, badge_text, title_text, subtitle_text=""):
    # Category badge
    if badge_text:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), Inches(0.35), Inches(3.4), Inches(0.32)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_BORDER_CYAN
        badge.line.fill.background()
        tf = badge.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = badge_text.upper()
        p.font.name = FONT_HEADING
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

    # Slide Title
    tb_title = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.72), Inches(12.13), Inches(0.55)
    )
    tf = tb_title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_HEADING
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Subtitle
    if subtitle_text:
        tb_sub = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.25), Inches(12.13), Inches(0.35)
        )
        tf = tb_sub.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = subtitle_text
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_CYAN_LIGHT

def add_footer(slide, slide_num, total_slides=18):
    tb_foot = slide.shapes.add_textbox(
        Inches(0.6), Inches(7.05), Inches(10.0), Inches(0.3)
    )
    tf = tb_foot.text_frame
    p = tf.paragraphs[0]
    p.text = "AI-Based Flood Detection & Mapping  |  Sen1Floods11 • Sentinel-1 SAR • PyTorch U-Net  |  Academic Minor Project"
    p.font.name = FONT_BODY
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED

    tb_num = slide.shapes.add_textbox(
        Inches(11.5), Inches(7.05), Inches(1.2), Inches(0.3)
    )
    tf = tb_num.text_frame
    p = tf.paragraphs[0]
    p.text = f"{slide_num} / {total_slides}"
    p.font.name = FONT_BODY
    p.font.size = Pt(9)
    p.font.color.rgb = C_MUTED
    p.alignment = PP_ALIGN.RIGHT

def create_card(slide, left, top, width, height, bg_color=C_CARD, border_color=C_BORDER, line_width=1):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(line_width)
    else:
        card.line.fill.background()
    return card

def add_card_header(slide, left, top, width, text, color=C_CYAN, size=13):
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_HEADING
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return tb

def add_speaker_notes(slide, notes):
    slide.notes_slide.notes_text_frame.text = notes.strip()

def build_presentation(assets_dir="/Users/supryo/Desktop/Ai-flood"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)
    create_card(s1, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3), bg_color=C_CARD_DARK, border_color=C_BORDER_CYAN, line_width=1.5)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(4.0), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_BORDER_CYAN
    badge.line.fill.background()
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "ACADEMIC MINOR PROJECT • AI & REMOTE SENSING"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    tb_t = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.8), Inches(1.4))
    tf = tb_t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AI-Based Flood Detection & Mapping Using Satellite Imagery"
    p.font.name = FONT_HEADING
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    tb_sub = s1.shapes.add_textbox(Inches(1.0), Inches(2.95), Inches(5.8), Inches(0.45))
    tf = tb_sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Sen1Floods11 • Sentinel-1 SAR Dual-Polarization • PyTorch U-Net"
    p.font.name = FONT_BODY
    p.font.size = Pt(13)
    p.font.color.rgb = C_CYAN_LIGHT

    create_card(s1, Inches(1.0), Inches(3.5), Inches(5.6), Inches(2.35), bg_color=C_CARD, border_color=C_BORDER)
    tb_hl = s1.shapes.add_textbox(Inches(1.15), Inches(3.6), Inches(5.3), Inches(2.15))
    tf = tb_hl.text_frame
    tf.word_wrap = True
    lines = [
        ("• Modality: ", "Sentinel-1 SAR C-Band (VV + VH Dual Polarization)"),
        ("• Architecture: ", "PyTorch U-Net with Group Normalization (G=8)"),
        ("• Benchmark: ", "Sen1Floods11 (11 Global Events across 6 Continents)"),
        ("• Evaluation: ", "100% Full Test Split (90/90 Chips, 20.5M Valid Pixels)"),
        ("• Test Results: ", "0.5548 IoU  |  0.7137 Dice/F1  |  93.34% Accuracy"),
        ("• Deployment: ", "Interactive Streamlit Dashboard + Custom GeoTIFF Upload")
    ]
    for i, (k, v) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = k
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN
        r2 = p.add_run()
        r2.text = v
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = C_TEXT_LIGHT

    tb_tech = s1.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(5.6), Inches(0.35))
    tf = tb_tech.text_frame
    p = tf.paragraphs[0]
    p.text = "TECH STACK: Python 3.11  •  PyTorch  •  U-Net  •  Streamlit  •  Rasterio  •  Torchvision"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_MUTED

    img1 = os.path.join(assets_dir, "outputs/visualizations/final_test/1_good_prediction_USA_905409.png")
    if os.path.exists(img1):
        s1.shapes.add_picture(img1, Inches(6.8), Inches(1.3), width=Inches(5.6))
        tb_cap = s1.shapes.add_textbox(Inches(6.8), Inches(6.1), Inches(5.6), Inches(0.3))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Verified Model Output: 5-Panel Flood Segmentation on USA Benchmark Chip (IoU: 0.9350)"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s1, """
Good morning/afternoon respected examiners and project guides.
Welcome to the presentation of our academic minor project: 'AI-Based Flood Detection and Mapping Using Satellite Imagery'.
In this work, we developed an audited deep learning semantic segmentation system using Sentinel-1 Synthetic Aperture Radar (SAR) imagery on the global Sen1Floods11 benchmark.
Our PyTorch U-Net incorporates Group Normalization with 8 groups and a domain-specific Masked BCE plus Dice loss, evaluated across 100% of the 90 official test chips (20.5 million valid pixels), achieving 0.5548 IoU and 0.7137 Dice.
We also built a live interactive Streamlit application for automated flood mapping and custom GeoTIFF upload. Let us examine the problem statement.
""")

    # =========================================================================
    # SLIDE 2: PROBLEM & MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Background & Motivation", "Disaster Response & The Optical Cloud Barrier", "Why automated radar-based flood mapping is critical for disaster assessment")
    add_footer(s2, 2)

    col_w = Inches(3.8)
    card_h = Inches(5.0)

    # Card 1
    create_card(s2, Inches(0.6), Inches(1.75), col_w, card_h)
    add_card_header(s2, Inches(0.8), Inches(1.9), col_w - Inches(0.4), "1. The Global Flood Crisis", color=C_AMBER, size=13)
    tb = s2.shapes.add_textbox(Inches(0.8), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    pts1 = [
        "• Most Frequent Natural Disaster: Floods account for over 40% of all disaster events worldwide.",
        "• Rapid Inundation: Floodwaters peak rapidly, isolating communities and submerging transport corridors.",
        "• Emergency Response Need: Disaster relief teams require rapid, accurate spatial maps of inundated zones.",
        "• Ground Survey Hazards: Field surveys are impeded by submerged infrastructure and hazardous conditions."
    ]
    for i, pt in enumerate(pts1):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(10)

    # Card 2
    create_card(s2, Inches(4.766), Inches(1.75), col_w, card_h)
    add_card_header(s2, Inches(4.966), Inches(1.9), col_w - Inches(0.4), "2. The Optical Sensor Obstruction", color=C_RED, size=13)
    tb = s2.shapes.add_textbox(Inches(4.966), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    pts2 = [
        "• Cloud Cover Obstruction: Active flood events are accompanied by severe storms, resulting in widespread cloud cover.",
        "• Optical Limitation: Sensors like Sentinel-2 and Landsat operate in visible/NIR spectra and cannot see through thick storm clouds.",
        "• Delayed Observation: Usable optical scenes are often unavailable until days or weeks after flood peaks.",
        "• Manual Interpretation Bottleneck: Manual digitization of large satellite swaths is too slow for emergency workflows."
    ]
    for i, pt in enumerate(pts2):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(10)

    # Card 3
    create_card(s2, Inches(8.933), Inches(1.75), col_w, card_h, border_color=C_BORDER_CYAN)
    add_card_header(s2, Inches(9.133), Inches(1.9), col_w - Inches(0.4), "3. The AI + SAR Solution", color=C_CYAN, size=13)
    tb = s2.shapes.add_textbox(Inches(9.133), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    pts3 = [
        "• All-Weather Penetration: Synthetic Aperture Radar (SAR) microwave signals penetrate rain, storm clouds, and operate day and night.",
        "• Radar Backscatter Contrast: Smooth standing water specularly reflects radar pulses away, appearing characteristically dark.",
        "• Deep Learning Automation: Custom U-Net model segments water extent at pixel level directly from dual-polarization SAR.",
        "• Fast Execution: Generates complete flood probability maps in under 50ms per 256×256 chip."
    ]
    for i, pt in enumerate(pts3):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(10)

    add_speaker_notes(s2, """
This slide outlines the fundamental challenge motivating this project.
Floods cause devastating social and economic losses. Timely disaster relief requires rapid mapping of inundated zones.
However, optical satellites like Sentinel-2 rely on visible sunlight and are severely obstructed by storm clouds during active flooding events.
Ground surveys are dangerous and hindered by submerged roads.
Synthetic Aperture Radar (SAR) provides the ideal solution because C-band microwaves penetrate cloud cover, rain, and darkness. Smooth floodwaters reflect radar pulses away, creating distinct contrast against surrounding terrain.
By pairing Sentinel-1 SAR with an optimized U-Net, our system automates pixel-level flood mapping in sub-50ms execution time.
""")

    # =========================================================================
    # SLIDE 3: OBJECTIVES
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Scope & Goals", "Project Objectives & Technical Deliverables", "Five core research and engineering milestones established for the system")
    add_footer(s3, 3)

    objectives = [
        ("1. Benchmark Ingestion & Zero-Leakage Audit", "Ingest and validate the global Sen1Floods11 benchmark across 11 flood events. Enforce strict mutual exclusivity across train, validation, and test splits with 0 overlapping chips.", C_CYAN),
        ("2. Dual-Polarization SAR Feature Extraction", "Exploit Sentinel-1 C-band SAR backscatter in both VV and VH channels, normalized from decibels to handle specular water reflection and land volume scattering.", C_GREEN),
        ("3. Optimized PyTorch U-Net Architecture", "Implement an efficient U-Net with Group Normalization (G=8) to ensure batch-size-independent normalization stability and spatial dropout (p=0.10) to mitigate overfitting.", C_CYAN_LIGHT),
        ("4. Masked Loss Formulation & Full Benchmark Evaluation", "Formulate a hybrid loss combining BCE and Dice loss calculated exclusively on valid pixels (ignoring -1 invalid labels). Rigorously evaluate across 100% of the 90 official test chips.", C_AMBER),
        ("5. Interactive Web Application & Geospatial Deployment", "Develop a lightweight, high-performance Streamlit dashboard supporting benchmark chip inspection, real-time threshold tuning, nominal flooded area estimation, and custom GeoTIFF upload.", C_WHITE)
    ]

    for i, (title, desc, accent) in enumerate(objectives):
        top_pos = Inches(1.75 + i * 1.0)
        create_card(s3, Inches(0.6), top_pos, Inches(12.133), Inches(0.9), bg_color=C_CARD, border_color=C_BORDER)

        p_badge = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), top_pos + Inches(0.18), Inches(0.55), Inches(0.55))
        p_badge.fill.solid()
        p_badge.fill.fore_color.rgb = accent
        p_badge.line.fill.background()
        tf = p_badge.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_BG
        p.alignment = PP_ALIGN.CENTER

        tb = s3.shapes.add_textbox(Inches(1.55), top_pos + Inches(0.1), Inches(11.0), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = accent

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(s3, """
Here are our five core project objectives:
First, to ingest and audit the global Sen1Floods11 benchmark dataset, verifying zero data leakage across splits.
Second, to exploit Sentinel-1 C-band SAR backscatter in dual polarization (VV + VH).
Third, to construct a customized PyTorch U-Net with Group Normalization (G=8), ensuring batch-size-independent normalization.
Fourth, to formulate a domain-specific Masked BCE plus Dice loss that ignores invalid -1 pixels and handles class imbalance, followed by full 90-chip test evaluation.
Fifth, to deploy an interactive Streamlit application with custom GeoTIFF upload and nominal flooded area estimation.
""")

    # =========================================================================
    # SLIDE 4: WHY SENTINEL-1 SAR (ENLARGED IMAGE: 50% OF SLIDE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Radar Remote Sensing", "Why Sentinel-1 Synthetic Aperture Radar (SAR)?", "Physical principles of microwave backscatter and dual-polarization discrimination")
    add_footer(s4, 4)

    # Left: Text card (5.4 inches)
    create_card(s4, Inches(0.6), Inches(1.75), Inches(5.4), Inches(5.0))
    add_card_header(s4, Inches(0.8), Inches(1.9), Inches(5.0), "Radar Backscatter Principles", color=C_CYAN, size=13)

    tb_sar = s4.shapes.add_textbox(Inches(0.8), Inches(2.35), Inches(5.0), Inches(4.2))
    tf = tb_sar.text_frame
    tf.word_wrap = True
    sar_pts = [
        ("• C-Band Active Radar (5.405 GHz): ", "Operates day and night independent of sunlight; penetrates clouds, rain, and haze."),
        ("• Specular Reflection on Water: ", "Smooth water acts like a mirror, reflecting radar pulses away from the antenna. Water appears characteristically dark (< -18 dB)."),
        ("• Diffuse Scattering on Land: ", "Rough soil, vegetation, and terrain scatter pulses diffusely back to the sensor, appearing brighter (-12 dB to -5 dB)."),
        ("• VV Polarization Channel: ", "Vertical transmit / Vertical receive. Highly sensitive to surface water roughness and open water boundaries."),
        ("• VH Polarization Channel: ", "Vertical transmit / Horizontal receive. Captures depolarized volume scattering from vegetation canopy.")
    ]
    for bp, tp in sar_pts:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = bp
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = tp
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    # Right: Large Dataset Verification Image (6.6 inches wide, 4.4 inches high)
    create_card(s4, Inches(6.2), Inches(1.75), Inches(6.533), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(s4, Inches(6.4), Inches(1.9), Inches(6.1), "Dual-Polarization Dataset Sample (Ghana)", color=C_GREEN, size=13)

    img_s4 = os.path.join(assets_dir, "outputs/visualizations/dataset_samples/Ghana_313799_verification.png")
    if os.path.exists(img_s4):
        s4.shapes.add_picture(img_s4, Inches(6.4), Inches(2.3), width=Inches(6.13))
        tb_cap = s4.shapes.add_textbox(Inches(6.4), Inches(6.35), Inches(6.13), Inches(0.35))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Sen1Floods11 Hand-Labeled Verification: S1 VV, S1 VH, S2 Optical, and Target Mask"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s4, """
Slide 4 details the physics of Synthetic Aperture Radar.
Sentinel-1 operates in C-band microwave spectrum at 5.405 GHz.
Because it is an active radar sensor, it emits its own pulses and captures the returned backscatter.
Smooth standing water specularly reflects radar energy away from the sensor, appearing distinctively dark, typically below -18 dB.
Rough terrain and vegetation scatter radar energy back diffusely, appearing brighter.
We use dual polarization:
The VV channel is vertically polarized and is sensitive to water surface roughness.
The VH cross-polarization channel captures volume scattering from vegetation, helping the model differentiate crops from inundated ground.
On the right, the enlarged verification panel demonstrates how the radar channels clearly delineate water bodies corresponding with the hand-labeled target mask.
""")

    # =========================================================================
    # SLIDE 5: DATASET SEN1FLOODS11 (ENLARGED BOLIVIA IMAGE)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Data Engineering & Audit", "Benchmark Dataset: Sen1Floods11", "Global distribution, split verification, and discrete ground truth label encoding")
    add_footer(s5, 5)

    # Left: Split and label info (5.4 inches)
    w_left = Inches(5.4)
    create_card(s5, Inches(0.6), Inches(1.75), w_left, Inches(2.4))
    add_card_header(s5, Inches(0.8), Inches(1.9), w_left - Inches(0.4), "Official Dataset Partitions", color=C_CYAN, size=13)

    tb_sp = s5.shapes.add_textbox(Inches(0.8), Inches(2.3), w_left - Inches(0.4), Inches(1.7))
    tf = tb_sp.text_frame
    splits = [
        ("• Training Split:", "252 chips", "(1,008 GeoTIFFs)"),
        ("• Validation Split:", "89 chips", "(356 GeoTIFFs)"),
        ("• Official Test Split:", "90 chips", "(360 GeoTIFFs)"),
        ("• Bolivia Holdout:", "15 chips", "(60 GeoTIFFs)")
    ]
    for s_name, count, detail in splits:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{s_name:<22} "
        r1.font.bold = True
        r1.font.color.rgb = C_CYAN_LIGHT
        r1.font.size = Pt(10.5)
        r2 = p.add_run()
        r2.text = f"{count:<10} "
        r2.font.bold = True
        r2.font.color.rgb = C_WHITE
        r2.font.size = Pt(10.5)
        r3 = p.add_run()
        r3.text = detail
        r3.font.color.rgb = C_MUTED
        r3.font.size = Pt(9.5)

    create_card(s5, Inches(0.6), Inches(4.3), w_left, Inches(2.45))
    add_card_header(s5, Inches(0.8), Inches(4.45), w_left - Inches(0.4), "Label Encoding & Audit", color=C_GREEN, size=13)

    tb_lbl = s5.shapes.add_textbox(Inches(0.8), Inches(4.85), w_left - Inches(0.4), Inches(1.8))
    tf = tb_lbl.text_frame
    lbl_pts = [
        ("• Zero Data Leakage: ", "Verified pairwise mutual exclusivity across splits — 0 overlapping chips."),
        ("• Label -1 (Invalid): ", "Missing data or sensor artifacts (masked out during loss calculation)."),
        ("• Label 0 (Dry Land): ", "Uninundated terrain, soil, roads, and non-flooded vegetation."),
        ("• Label 1 (Water/Flood): ", "Inundated floodplains and water bodies (12.51% of valid test pixels)."),
        ("• Historical Baseline: ", "JRC 30-year permanence raster provides historical surface water baseline.")
    ]
    for bp, tp in lbl_pts:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = bp
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_GREEN if "Zero" in bp else (C_AMBER if "-1" in bp else C_WHITE)
        r2 = p.add_run()
        r2.text = tp
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT

    # Right: Large Bolivia Verification Image (6.5 inches wide)
    create_card(s5, Inches(6.2), Inches(1.75), Inches(6.533), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(s5, Inches(6.4), Inches(1.9), Inches(6.1), "Bolivia Holdout Verification Sample", color=C_WHITE, size=13)

    img_s5 = os.path.join(assets_dir, "outputs/visualizations/dataset_samples/Bolivia_103757_verification.png")
    if os.path.exists(img_s5):
        s5.shapes.add_picture(img_s5, Inches(6.4), Inches(2.3), width=Inches(6.13))
        tb_cap = s5.shapes.add_textbox(Inches(6.4), Inches(6.35), Inches(6.13), Inches(0.35))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Bolivia Holdout Event: S1 VV, S1 VH, S2 Optical, and Discrete Ground Truth Target"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s5, """
Slide 5 introduces the Sen1Floods11 benchmark dataset, spanning 11 catastrophic flood events across 6 continents.
The dataset is cleanly split into 252 training chips, 89 validation chips, 90 official test chips, and 15 Bolivia holdout chips.
We programmatically audited the dataset splits to verify zero overlapping chips between training and test partitions.
The ground truth masks use discrete integer encoding:
-1 for invalid pixels (which are ignored during training), 0 for dry land, and 1 for floodwater.
We also clarify that Sen1Floods11 does not provide separate pre-flood SAR acquisitions; the JRC Global Surface Water 30-year permanence raster is used as the historical water baseline.
""")

    # =========================================================================
    # SLIDE 6: PREPROCESSING PIPELINE
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Pipeline Engineering", "Data Preprocessing & Tensor Pipeline", "End-to-end normalization, strided spatial subsampling, and synchronized augmentation")
    add_footer(s6, 6)

    step_w = Inches(2.25)
    step_h = Inches(4.8)
    gap = Inches(0.2)
    start_x = Inches(0.6)

    steps = [
        ("Step 1", "Decibel Clipping", C_CYAN, [
            "• Input: Float32 GeoTIFFs.",
            "• Range: [-35 dB, +5 dB].",
            "• Clips extreme specular outliers below -35 dB and corner reflectors above +5 dB.",
            "• Suppresses radar flare artifacts."
        ]),
        ("Step 2", "Min-Max Scaling", C_CYAN_LIGHT, [
            "• Normalized Range: [0.0, 1.0].",
            "• Formula: (dB - (-35)) / 40.0.",
            "• Applied identically to VV and VH channels.",
            "• Ensures stable numerical gradient updates."
        ]),
        ("Step 3", "Strided Subsampling", C_GREEN, [
            "• Target Resolution: 256 × 256.",
            "• Method: Slicing ([::stride, ::stride]).",
            "• Strictly preserves discrete integer labels (-1, 0, 1).",
            "• Avoids interpolation blurring."
        ]),
        ("Step 4", "Synchronized Augment", C_AMBER, [
            "• Random Horizontal Flip (p=0.5).",
            "• Random Vertical Flip (p=0.5).",
            "• Random 90° Rotations (p=0.5).",
            "• Applied identically to SAR channels and masks."
        ]),
        ("Step 5", "Tensor Assembly", C_WHITE, [
            "• Tensor Shape: (2, 256, 256).",
            "• Channel 0: Normalized VV.",
            "• Channel 1: Normalized VH.",
            "• Target: (1, 256, 256) LongTensor.",
            "• Formats data for U-Net input."
        ])
    ]

    for i, (s_num, s_title, s_color, s_bullets) in enumerate(steps):
        x = start_x + i * (step_w + gap)
        create_card(s6, x, Inches(1.85), step_w, step_h)

        badge_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), Inches(2.0), step_w - Inches(0.4), Inches(0.32))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = s_color
        badge_box.line.fill.background()
        tf = badge_box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = s_num.upper()
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_BG if s_color != C_CARD else C_WHITE
        p.alignment = PP_ALIGN.CENTER

        tb = s6.shapes.add_textbox(x + Inches(0.15), Inches(2.4), step_w - Inches(0.3), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER

        tb_b = s6.shapes.add_textbox(x + Inches(0.15), Inches(2.9), step_w - Inches(0.3), Inches(3.6))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        for j, b_text in enumerate(s_bullets):
            p = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            p.text = b_text
            p.font.size = Pt(10)
            p.font.color.rgb = C_TEXT_LIGHT
            p.space_after = Pt(6)

    add_speaker_notes(s6, """
Slide 6 illustrates our data preprocessing pipeline.
Step 1 clips raw decibel backscatter to the physical range of -35 to +5 dB.
Step 2 normalizes this range to [0.0, 1.0] for neural network numerical stability.
Step 3 is critical: we use strided slicing to resize satellite chips to 256 by 256 without applying interpolation algorithms, preserving discrete mask values.
Step 4 applies synchronized geometric augmentations (flips and 90-degree rotations).
Step 5 stacks the normalized VV and VH channels into a 2-channel PyTorch tensor ready for the U-Net.
""")

    # =========================================================================
    # SLIDE 7: U-NET ARCHITECTURE (VISUAL ENCODER-DECODER WITH ARROWS)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Model Design", "Customized PyTorch U-Net Architecture", "Recognizable encoder-decoder network with Group Normalization (G=8) and skip connections")
    add_footer(s7, 7)

    # Main Architecture Diagram Card (Left, 7.8 inches)
    create_card(s7, Inches(0.6), Inches(1.75), Inches(8.0), Inches(5.0))
    add_card_header(s7, Inches(0.8), Inches(1.9), Inches(7.6), "Encoder-Decoder Flow with Skip Connections", color=C_CYAN, size=13)

    # Input Box
    in_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.35), Inches(7.2), Inches(0.38))
    in_box.fill.solid()
    in_box.fill.fore_color.rgb = C_CARD_DARK
    in_box.line.color.rgb = C_CYAN
    tf = in_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "INPUT: 2 × 256 × 256 Tensor (Channel 0: S1-VV, Channel 1: S1-VH)"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_CYAN
    p.alignment = PP_ALIGN.CENTER

    # 4 Encoder vs 4 Decoder Level Blocks
    levels = [
        ("Encoder 1 (32 ch, 128²)", "Decoder 1 (32 ch, 256²)", Inches(2.85)),
        ("Encoder 2 (64 ch, 64²)",  "Decoder 2 (64 ch, 128²)", Inches(3.30)),
        ("Encoder 3 (128 ch, 32²)", "Decoder 3 (128 ch, 64²)",  Inches(3.75)),
        ("Encoder 4 (256 ch, 16²)", "Decoder 4 (256 ch, 32²)",  Inches(4.20))
    ]

    for enc_txt, dec_txt, y_pos in levels:
        # Encoder Block
        enc_b = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y_pos, Inches(2.6), Inches(0.36))
        enc_b.fill.solid()
        enc_b.fill.fore_color.rgb = C_CARD_DARK
        enc_b.line.color.rgb = C_BORDER_CYAN
        tf = enc_b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = enc_txt
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_CYAN_LIGHT
        p.alignment = PP_ALIGN.CENTER

        # Skip Connection Arrow Shape
        skip_arrow = s7.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.7), y_pos + Inches(0.08), Inches(1.8), Inches(0.2))
        skip_arrow.fill.solid()
        skip_arrow.fill.fore_color.rgb = C_GREEN
        skip_arrow.line.fill.background()
        tf = skip_arrow.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "Skip Concatenation"
        p.font.size = Pt(7.5)
        p.font.bold = True
        p.font.color.rgb = C_BG
        p.alignment = PP_ALIGN.CENTER

        # Decoder Block
        dec_b = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), y_pos, Inches(2.6), Inches(0.36))
        dec_b.fill.solid()
        dec_b.fill.fore_color.rgb = C_CARD_DARK
        dec_b.line.color.rgb = C_GREEN
        tf = dec_b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = dec_txt
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_GREEN
        p.alignment = PP_ALIGN.CENTER

    # Bottleneck Box
    btn_b = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.70), Inches(7.2), Inches(0.42))
    btn_b.fill.solid()
    btn_b.fill.fore_color.rgb = C_PURPLE
    btn_b.line.fill.background()
    tf = btn_b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "★ BOTTLENECK: Conv3x3 (256 → 512) + Spatial Dropout (p=0.10) [512 × 16 × 16]"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Output Box
    out_b = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.25), Inches(7.2), Inches(0.38))
    out_b.fill.solid()
    out_b.fill.fore_color.rgb = C_GREEN
    out_b.line.fill.background()
    tf = out_b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "OUTPUT: 1×1 Conv (32 → 1) → Single-Channel Logits (1 × 256 × 256)"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_BG
    p.alignment = PP_ALIGN.CENTER

    # Note on GroupNorm
    tb_gn = s7.shapes.add_textbox(Inches(1.0), Inches(5.75), Inches(7.2), Inches(0.85))
    tf = tb_gn.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "• Block Design: Each stage utilizes DoubleConv3x3 + GroupNorm (G=8) + GELU activation.\n• Upsampling: Bilinear interpolation avoids checkerboard deconvolution artifacts."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_TEXT_LIGHT

    # Right: Specifications Card (3.9 inches)
    create_card(s7, Inches(8.8), Inches(1.75), Inches(3.933), Inches(5.0))
    add_card_header(s7, Inches(9.0), Inches(1.9), Inches(3.5), "Architecture Specifications", color=C_GREEN, size=13)

    tb_spec = s7.shapes.add_textbox(Inches(9.0), Inches(2.35), Inches(3.5), Inches(4.2))
    tf_spec = tb_spec.text_frame
    tf_spec.word_wrap = True
    specs = [
        ("GroupNorm (G=8) vs BatchNorm:", "BatchNorm relies on batch statistics. With small batch size (B=2), BatchNorm exhibits high statistical variance. GroupNorm does not depend on batch-level statistics, operating per image across 8 channel groups."),
        ("Bilinear Upsampling:", "Replaces transposed convolution to ensure artifact-free reconstruction along narrow waterways."),
        ("Skip Connections:", "Preserves fine spatial boundaries by concatenating encoder feature maps directly to decoder stages."),
        ("Parameters: 7,760,257", "Balanced capacity providing representational expressiveness while maintaining <50ms inference latency.")
    ]
    for h, b in specs:
        p = tf_spec.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {h}\n  "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    add_speaker_notes(s7, """
Slide 7 illustrates our customized PyTorch U-Net architecture.
The network follows an encoder-decoder topology with skip connections:
The encoder contracts features across 4 stages (32, 64, 128, 256 channels), followed by a 512-channel bottleneck with spatial dropout (p=0.10).
The decoder expands feature maps back to 256x256 resolution using bilinear upsampling and skip concatenation.
A central architectural decision is Group Normalization with 8 groups:
Because satellite imagery batch size is limited (B=2), Batch Normalization produces unstable running statistics. GroupNorm does not depend on batch-level statistics, ensuring stable training.
Total trainable parameters: exactly 7,760,257.
""")

    # =========================================================================
    # SLIDE 8: LOSS & TRAINING (ENLARGED LOSS CURVE)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Optimization & Convergence", "Loss Formulation & Training Configuration", "Hybrid masked loss overcoming class imbalance and verified convergence curves")
    add_footer(s8, 8)

    # Left: Loss Formula Card (5.4 inches)
    w_l = Inches(5.4)
    create_card(s8, Inches(0.6), Inches(1.75), w_l, Inches(5.0))
    add_card_header(s8, Inches(0.8), Inches(1.9), w_l - Inches(0.4), "Masked BCE + Dice Loss Formulation", color=C_CYAN, size=13)

    tb_loss = s8.shapes.add_textbox(Inches(0.8), Inches(2.35), w_l - Inches(0.4), Inches(4.2))
    tf = tb_loss.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Loss Function (Evaluated on Valid Pixels y ≠ -1):"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = "ℒ_total = 0.5 × ℒ_BCE + 0.5 × ℒ_Dice"
    p2.font.name = "Courier New"
    p2.font.bold = True
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_CYAN
    p2.space_before = Pt(4)
    p2.space_after = Pt(8)

    loss_pts = [
        ("1. Binary Cross-Entropy (BCE):", "Calculates pixel-wise logarithmic loss for smooth probability convergence."),
        ("2. Soft Dice Loss:", "Directly optimizes region overlap, mitigating severe foreground class imbalance (12.5% water vs. 87.5% land)."),
        ("3. Dynamic Masking:", "Pixels labeled -1 are masked out prior to loss computation, preventing corrupted gradient propagation."),
        ("4. Training Config:", "AdamW (lr = 0.0005, weight decay = 1e-4), Cosine Annealing, 5 epochs in ~15.2s on Apple Silicon MPS.")
    ]
    for h, b in loss_pts:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{h} "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    # Right: Large Training Loss Curve (6.5 inches wide)
    create_card(s8, Inches(6.2), Inches(1.75), Inches(6.533), Inches(5.0), bg_color=C_CARD_DARK)
    add_card_header(s8, Inches(6.4), Inches(1.9), Inches(6.1), "Training & Validation Loss Convergence", color=C_GREEN, size=13)

    img_s8 = os.path.join(assets_dir, "outputs/training/loss_curves/s1_loss_curve.png")
    if os.path.exists(img_s8):
        s8.shapes.add_picture(img_s8, Inches(6.4), Inches(2.3), width=Inches(6.13))
        tb_cap = s8.shapes.add_textbox(Inches(6.4), Inches(6.35), Inches(6.13), Inches(0.35))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Verified Loss Curves: Steady convergence across 5 epochs with Cosine Annealing"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s8, """
Slide 8 covers our loss formulation and training setup.
Due to severe class imbalance—with water occupying only 12.5% of pixels—we combined Binary Cross-Entropy with Soft Dice loss:
0.5 BCE plus 0.5 Dice loss.
BCE provides smooth pixel classification convergence, while Dice loss directly optimizes spatial region overlap.
Invalid -1 pixels are masked out dynamically.
We trained for 5 epochs using AdamW with learning rate 0.0005 and Cosine Annealing.
On the right, the enlarged training curve confirms smooth, monotonic decrease in validation loss without divergence.
""")

    # =========================================================================
    # SLIDE 9: EVALUATION METRICS
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Quantitative Assessment", "Segmentation Evaluation Metrics & Formulations", "Formulations evaluated strictly over valid pixels V = {i | y_i != -1}")
    add_footer(s9, 9)

    c_w = Inches(5.8)
    c_h = Inches(1.9)
    metrics_list = [
        ("Intersection over Union (IoU / Jaccard Index)", "IoU = TP / (TP + FP + FN)", "Measures geometric overlap between predicted and ground-truth flood masks. Primary benchmark standard for remote sensing segmentation.", C_CYAN, Inches(0.6), Inches(1.75)),
        ("Dice Coefficient / F1 Score", "Dice = 2·TP / (2·TP + FP + FN) = F1", "Harmonic mean of precision and recall. Directly penalizes both false alarms (FP) and missed floodwaters (FN).", C_GREEN, Inches(6.9), Inches(1.75)),
        ("Precision (Positive Predictive Value)", "Precision = TP / (TP + FP)", "Proportion of predicted flood pixels that are true water. Crucial for reducing false-positive flood alarms.", C_CYAN_LIGHT, Inches(0.6), Inches(3.8)),
        ("Recall (Sensitivity / Detection Rate)", "Recall = TP / (TP + FN)", "Proportion of actual floodwaters successfully identified. Essential for disaster coverage.", C_AMBER, Inches(6.9), Inches(3.8))
    ]

    for title, formula, desc, color, x, y in metrics_list:
        create_card(s9, x, y, c_w, c_h)
        add_card_header(s9, x + Inches(0.2), y + Inches(0.12), c_w - Inches(0.4), title, color=color, size=11.5)

        tb = s9.shapes.add_textbox(x + Inches(0.2), y + Inches(0.45), c_w - Inches(0.4), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = formula
        p.font.name = "Courier New"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT

    create_card(s9, Inches(0.6), Inches(5.85), Inches(12.1), Inches(0.95), bg_color=C_CARD_DARK, border_color=C_BORDER_CYAN)
    tb_bot = s9.shapes.add_textbox(Inches(0.8), Inches(5.92), Inches(11.7), Inches(0.8))
    tf_b = tb_bot.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "CRITICAL VIVA INSIGHT — Why Overall Accuracy is Deceptive in Geospatial Tasks:"
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_AMBER
    p2 = tf_b.add_paragraph()
    p2.text = "Because 87.5% of valid test pixels are dry land, a trivial model predicting 100% dry land would obtain 87.5% accuracy while failing to detect any floodwater. Thus, IoU (0.5548) and Dice (0.7137) serve as the true benchmarks of segmentation capability."
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(s9, """
Slide 9 defines our evaluation metrics.
All metrics are evaluated exclusively over valid pixels, strictly ignoring -1 invalid labels:
IoU: True Positives over True Positives plus False Positives plus False Negatives.
Dice Score: Harmonic mean of precision and recall.
Precision: True Positives over all predicted positives—reducing false alarms.
Recall: True Positives over all true flood pixels.
Regarding accuracy: Because 87.5% of the benchmark is dry land, accuracy alone is misleading. IoU and Dice are the primary standards for rigorous semantic segmentation evaluation.
""")

    # =========================================================================
    # SLIDE 10: OFFICIAL TEST RESULTS (ENLARGED METRIC GRAPH)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "Quantitative Benchmark", "Official Complete Test Results (90 / 90 Chips)", "Verified performance evaluated across 100% of the official Sen1Floods11 test split")
    add_footer(s10, 10)

    # 5 KPI Cards
    kpis = [
        ("TEST IoU", "0.5548", "Jaccard Index", C_CYAN),
        ("DICE / F1", "0.7137", "Overlap Coeff", C_GREEN),
        ("PRECISION", "0.7717", "77.17% Water True", C_CYAN_LIGHT),
        ("RECALL", "0.6637", "66.37% Water Found", C_AMBER),
        ("ACCURACY", "0.9334", "93.34% Pixel Acc", C_WHITE)
    ]
    card_w = Inches(2.25)
    for i, (label, val, sub, col) in enumerate(kpis):
        x = Inches(0.6 + i * 2.47)
        create_card(s10, x, Inches(1.75), card_w, Inches(1.45), bg_color=C_CARD, border_color=C_BORDER)

        tb = s10.shapes.add_textbox(x + Inches(0.1), Inches(1.85), card_w - Inches(0.2), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = label
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.alignment = PP_ALIGN.CENTER

        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.alignment = PP_ALIGN.CENTER

    # Bottom Left: Summary Table (5.4 inches)
    w_l10 = Inches(5.4)
    create_card(s10, Inches(0.6), Inches(3.35), w_l10, Inches(3.4))
    add_card_header(s10, Inches(0.8), Inches(3.5), w_l10 - Inches(0.4), "Full Benchmark Evaluation Scope", color=C_CYAN, size=13)

    tb_scope = s10.shapes.add_textbox(Inches(0.8), Inches(3.9), w_l10 - Inches(0.4), Inches(2.7))
    tf = tb_scope.text_frame
    scope_rows = [
        ("• Test Chips Evaluated:", "90 / 90 (100% complete)"),
        ("• Total Valid Pixels:", "20,517,367 pixels"),
        ("• Ground Truth Water:", "2,566,101 pixels (12.51%)"),
        ("• Ground Truth Land:", "17,951,266 pixels (87.49%)"),
        ("• Test Masked Loss:", "0.3918"),
        ("• Benchmark Precision:", "shows 77.17% precision on official set")
    ]
    for k, v in scope_rows:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{k:<24} "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = v
        r2.font.bold = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_WHITE

    # Bottom Right: Large Metric Progression Curve (6.5 inches wide)
    create_card(s10, Inches(6.2), Inches(3.35), Inches(6.533), Inches(3.4), bg_color=C_CARD_DARK)
    add_card_header(s10, Inches(6.4), Inches(3.5), Inches(6.1), "Validation Metric Progression (IoU & Dice)", color=C_GREEN, size=13)

    img_s10 = os.path.join(assets_dir, "outputs/training/metric_curves/s1_metric_curve.png")
    if os.path.exists(img_s10):
        s10.shapes.add_picture(img_s10, Inches(6.4), Inches(3.85), width=Inches(6.13))

    add_speaker_notes(s10, """
Slide 10 presents our official complete test results.
We evaluated all 90 chips in the official Sen1Floods11 test split without exclusions, covering 20,517,367 valid pixels.
Verified test results:
Test IoU: 0.5548
Test Dice / F1 Score: 0.7137
Test Precision: 0.7717
Test Recall: 0.6637
Overall Accuracy: 0.9334
The model demonstrates 77.17% precision on the official benchmark.
On the right, the enlarged validation metric progression confirms that IoU and Dice plateaued stably by epoch 5.
""")

    # =========================================================================
    # SLIDE 11: CONFUSION MATRIX
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11)
    add_header(s11, "Error Diagnostics", "Pixel-Level Test Confusion Matrix", "Complete 20,517,367 pixel breakdown and error taxonomy across 90 test chips")
    add_footer(s11, 11)

    # Left: 2x2 Grid (5.6 inches)
    create_card(s11, Inches(0.6), Inches(1.75), Inches(5.6), Inches(5.0))
    add_card_header(s11, Inches(0.8), Inches(1.9), Inches(5.2), "Official Test Set Confusion Matrix", color=C_CYAN, size=13)

    # TP
    tp_b = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(2.35), Inches(2.45), Inches(1.3))
    tp_b.fill.solid()
    tp_b.fill.fore_color.rgb = C_GREEN_DARK
    tp_b.line.color.rgb = C_GREEN
    tf = tp_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRUE POSITIVES (TP)\n1,703,186 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Water correctly detected (66.4%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = C_GREEN
    p2.alignment = PP_ALIGN.CENTER

    # FP
    fp_b = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.55), Inches(2.35), Inches(2.45), Inches(1.3))
    fp_b.fill.solid()
    fp_b.fill.fore_color.rgb = C_RED_DARK
    fp_b.line.color.rgb = C_RED
    tf = fp_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FALSE POSITIVES (FP)\n503,818 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Dry land flagged as flood (2.8%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = RGBColor(252, 165, 165)
    p2.alignment = PP_ALIGN.CENTER

    # FN
    fn_b = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(3.8), Inches(2.45), Inches(1.3))
    fn_b.fill.solid()
    fn_b.fill.fore_color.rgb = RGBColor(120, 53, 15)
    fn_b.line.color.rgb = C_AMBER
    tf = fn_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FALSE NEGATIVES (FN)\n862,915 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Floodwater missed (33.6%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = C_AMBER
    p2.alignment = PP_ALIGN.CENTER

    # TN
    tn_b = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.55), Inches(3.8), Inches(2.45), Inches(1.3))
    tn_b.fill.solid()
    tn_b.fill.fore_color.rgb = C_CARD_DARK
    tn_b.line.color.rgb = C_BORDER
    tf = tn_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TRUE NEGATIVES (TN)\n17,447,448 Pixels"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = "Land correctly identified (97.2%)"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = C_MUTED
    p2.alignment = PP_ALIGN.CENTER

    tb_tot = s11.shapes.add_textbox(Inches(0.85), Inches(5.25), Inches(5.1), Inches(1.3))
    tf = tb_tot.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "• Total Evaluated Valid Pixels: 20,517,367 (100% of 90 test chips)\n• Ground Truth Water: 2,566,101 pixels (12.51%)\n• Ground Truth Land: 17,951,266 pixels (87.49%)"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_TEXT_LIGHT

    # Right: Error Diagnostics (6.3 inches)
    create_card(s11, Inches(6.4), Inches(1.75), Inches(6.333), Inches(5.0))
    add_card_header(s11, Inches(6.6), Inches(1.9), Inches(5.9), "Physical Radar Error Taxonomy", color=C_AMBER, size=13)

    tb_err = s11.shapes.add_textbox(Inches(6.6), Inches(2.35), Inches(5.9), Inches(4.2))
    tf = tb_err.text_frame
    tf.word_wrap = True
    err_pts = [
        ("1. False Positives (FP = 503,818 pixels):", [
            "• Smooth, flat dry surfaces (e.g. airport runways, dry playas, calm sand).",
            "• Mirror-like terrain produces specular reflection mimicking open water.",
            "• The model reduces the proportion of false-positive flood pixels in benchmark predictions using VH cross-polarization texture."
        ]),
        ("2. False Negatives (FN = 862,915 pixels):", [
            "• Flooded vegetation causing corner-reflector 'double-bounce' backscatter.",
            "• Narrow drainage ditches narrower than Sentinel-1's spatial resolution.",
            "• Surface roughening from high winds altering specular reflection."
        ]),
        ("3. Precision Impact:", [
            "• 77.17% Precision demonstrates low false-positive rate on land.",
            "• 66.37% Recall captures continuous open-water floodplains reliably."
        ])
    ]
    for head, bullets in err_pts:
        p = tf.add_paragraph()
        p.text = head
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_CYAN_LIGHT
        p.space_before = Pt(4)
        for b in bullets:
            p2 = tf.add_paragraph()
            p2.text = b
            p2.font.size = Pt(9.5)
            p2.font.color.rgb = C_TEXT_LIGHT

    add_speaker_notes(s11, """
Slide 11 breaks down our pixel-level confusion matrix across all 20.5 million valid test pixels:
True Positives: 1,703,186 pixels
False Positives: 503,818 pixels
False Negatives: 862,915 pixels
True Negatives: 17,447,448 pixels.
Analyzing radar physics explains these error distributions:
False Positives occur primarily on smooth dry land like airport runways that mirror radar pulses away, resembling calm water.
False Negatives occur when floodwaters submerge dense forest canopies, where tree trunks create double-bounce backscatter that appears bright on radar.
Understanding these radar phenomena helps us interpret model performance accurately.
""")

    # =========================================================================
    # SLIDE 12: GOOD PREDICTION — USA_905409 (DEDICATED FULL-WIDTH SLIDE)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12)
    add_header(s12, "Qualitative Case 1 (Good)", "Benchmark Test Case: USA_905409 (High-Contrast Inundation)", "Representative good prediction on open agricultural floodplains with high specular contrast")
    add_footer(s12, 12)

    # Top Metrics Banner Card
    create_card(s12, Inches(0.6), Inches(1.75), Inches(12.133), Inches(0.95), bg_color=C_CARD, border_color=C_GREEN)
    tb_m12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.82), Inches(11.7), Inches(0.8))
    tf = tb_m12.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CATEGORY: GOOD PREDICTION  |  EVENT: USA (FLOODPLAINS)  |  VALID PIXELS: 262,099"
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = C_GREEN

    p2 = tf.add_paragraph()
    r1 = p2.add_run(); r1.text = "Test IoU: "; r1.font.bold = True; r1.font.size = Pt(13); r1.font.color.rgb = C_CYAN
    r2 = p2.add_run(); r2.text = "0.9350    "; r2.font.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = C_WHITE
    r3 = p2.add_run(); r3.text = "Dice / F1: "; r3.font.bold = True; r3.font.size = Pt(13); r3.font.color.rgb = C_CYAN
    r4 = p2.add_run(); r4.text = "0.9664    "; r4.font.bold = True; r4.font.size = Pt(14); r4.font.color.rgb = C_WHITE
    r5 = p2.add_run(); r5.text = "Precision: "; r5.font.bold = True; r5.font.size = Pt(13); r5.font.color.rgb = C_CYAN
    r6 = p2.add_run(); r6.text = "0.9720    "; r6.font.bold = True; r6.font.size = Pt(14); r6.font.color.rgb = C_WHITE
    r7 = p2.add_run(); r7.text = "Recall: "; r7.font.bold = True; r7.font.size = Pt(13); r7.font.color.rgb = C_CYAN
    r8 = p2.add_run(); r8.text = "0.9609    "; r8.font.bold = True; r8.font.size = Pt(14); r8.font.color.rgb = C_WHITE
    r9 = p2.add_run(); r9.text = "Water Pixels: "; r9.font.bold = True; r9.font.size = Pt(11); r9.font.color.rgb = C_MUTED
    r10 = p2.add_run(); r10.text = "44,467"; r10.font.bold = True; r10.font.size = Pt(11); r10.font.color.rgb = C_TEXT_LIGHT

    # Large Prediction Figure (Occupies 70% width/height)
    create_card(s12, Inches(0.6), Inches(2.85), Inches(12.133), Inches(3.95), bg_color=C_CARD_DARK)
    img_s12 = os.path.join(assets_dir, "outputs/visualizations/final_test/1_good_prediction_USA_905409.png")
    if os.path.exists(img_s12):
        s12.shapes.add_picture(img_s12, Inches(0.8), Inches(3.0), width=Inches(11.73))
        tb_cap = s12.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.73), Inches(0.9))
        tf = tb_cap.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "Panels (L to R): 1. Input SAR False-Color (VV, VH, Ratio) | 2. Ground Truth Mask | 3. Predicted Flood Mask | 4. Model Overlay | 5. Flood Probability Map"
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_CYAN_LIGHT
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = "Analysis: The open agricultural floodplain provides sharp specular radar contrast, allowing U-Net to recover crisp flood boundaries with near-perfect 0.9350 IoU."
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s12, """
Slide 12 presents our first dedicated qualitative case study: USA_905409, representing a high-performing benchmark prediction.
The model achieved an IoU of 0.9350, Dice of 0.9664, Precision of 0.9720, and Recall of 0.9609.
Looking at the enlarged 5-panel visualization:
Panel 1 shows the SAR false-color composite.
Panel 2 is the ground truth mask.
Panel 3 is our model's binary prediction.
Panel 4 displays the color overlay.
Panel 5 displays the continuous Flood Probability Map.
Because open agricultural floodplains produce clean specular reflection, the U-Net accurately recovers flood boundaries with minimal boundary blur.
""")

    # =========================================================================
    # SLIDE 13: AVERAGE PREDICTION — PAKISTAN_694942 (DEDICATED FULL-WIDTH SLIDE)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13)
    add_header(s13, "Qualitative Case 2 (Average)", "Benchmark Test Case: Pakistan_694942 (Complex Boundaries)", "Representative average prediction on agrarian fields with mud-water transition zones")
    add_footer(s13, 13)

    create_card(s13, Inches(0.6), Inches(1.75), Inches(12.133), Inches(0.95), bg_color=C_CARD, border_color=C_CYAN)
    tb_m13 = s13.shapes.add_textbox(Inches(0.8), Inches(1.82), Inches(11.7), Inches(0.8))
    tf = tb_m13.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CATEGORY: AVERAGE PREDICTION  |  EVENT: PAKISTAN (AGRICULTURAL)  |  VALID PIXELS: 191,444"
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = C_CYAN

    p2 = tf.add_paragraph()
    r1 = p2.add_run(); r1.text = "Test IoU: "; r1.font.bold = True; r1.font.size = Pt(13); r1.font.color.rgb = C_CYAN
    r2 = p2.add_run(); r2.text = "0.3829    "; r2.font.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = C_WHITE
    r3 = p2.add_run(); r3.text = "Dice / F1: "; r3.font.bold = True; r3.font.size = Pt(13); r3.font.color.rgb = C_CYAN
    r4 = p2.add_run(); r4.text = "0.5538    "; r4.font.bold = True; r4.font.size = Pt(14); r4.font.color.rgb = C_WHITE
    r5 = p2.add_run(); r5.text = "Precision: "; r5.font.bold = True; r5.font.size = Pt(13); r5.font.color.rgb = C_CYAN
    r6 = p2.add_run(); r6.text = "0.4152    "; r6.font.bold = True; r6.font.size = Pt(14); r6.font.color.rgb = C_WHITE
    r7 = p2.add_run(); r7.text = "Recall: "; r7.font.bold = True; r7.font.size = Pt(13); r7.font.color.rgb = C_CYAN
    r8 = p2.add_run(); r8.text = "0.8313    "; r8.font.bold = True; r8.font.size = Pt(14); r8.font.color.rgb = C_WHITE
    r9 = p2.add_run(); r9.text = "Water Pixels: "; r9.font.bold = True; r9.font.size = Pt(11); r9.font.color.rgb = C_MUTED
    r10 = p2.add_run(); r10.text = "18,481"; r10.font.bold = True; r10.font.size = Pt(11); r10.font.color.rgb = C_TEXT_LIGHT

    create_card(s13, Inches(0.6), Inches(2.85), Inches(12.133), Inches(3.95), bg_color=C_CARD_DARK)
    img_s13 = os.path.join(assets_dir, "outputs/visualizations/final_test/2_average_prediction_Pakistan_694942.png")
    if os.path.exists(img_s13):
        s13.shapes.add_picture(img_s13, Inches(0.8), Inches(3.0), width=Inches(11.73))
        tb_cap = s13.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.73), Inches(0.9))
        tf = tb_cap.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "Panels (L to R): 1. Input SAR False-Color (VV, VH, Ratio) | 2. Ground Truth Mask | 3. Predicted Flood Mask | 4. Model Overlay | 5. Flood Probability Map"
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_CYAN_LIGHT
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = "Analysis: The model achieves high recall (83.13%), capturing major inundated zones, but exhibits false-positive boundary expansion along saturated mud-water transition boundaries."
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s13, """
Slide 13 shows our second case study: Pakistan_694942, representing an average-performance test chip.
The model achieved an IoU of 0.3829, Dice of 0.5538, and a high Recall of 0.8313.
In this scene, intense flooding submerged agricultural parcels and irrigation canals.
The high recall confirms that 83.1% of true floodwaters were captured.
However, precision is 41.52% due to false positives along saturated mudflats and wet soil, where soil moisture lowers radar backscatter and mimics shallow standing water.
The Flood Probability Map in Panel 5 clearly reflects intermediate probability values in these boundary zones.
""")

    # =========================================================================
    # SLIDE 14: DIFFICULT PREDICTION — SRI-LANKA_450918 (DEDICATED FULL-WIDTH SLIDE)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14)
    add_header(s14, "Qualitative Case 3 (Difficult)", "Benchmark Test Case: Sri-Lanka_450918 (Canopy Attenuation)", "Transparent evaluation of radar double-bounce limitations in dense tropical vegetation")
    add_footer(s14, 14)

    create_card(s14, Inches(0.6), Inches(1.75), Inches(12.133), Inches(0.95), bg_color=C_CARD, border_color=C_AMBER)
    tb_m14 = s14.shapes.add_textbox(Inches(0.8), Inches(1.82), Inches(11.7), Inches(0.8))
    tf = tb_m14.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CATEGORY: DIFFICULT PREDICTION (PHYSICAL RADAR LIMITATION)  |  EVENT: SRI LANKA  |  VALID PIXELS: 99,126"
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = C_AMBER

    p2 = tf.add_paragraph()
    r1 = p2.add_run(); r1.text = "Test IoU: "; r1.font.bold = True; r1.font.size = Pt(13); r1.font.color.rgb = C_AMBER
    r2 = p2.add_run(); r2.text = "0.0086    "; r2.font.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = C_WHITE
    r3 = p2.add_run(); r3.text = "Dice / F1: "; r3.font.bold = True; r3.font.size = Pt(13); r3.font.color.rgb = C_AMBER
    r4 = p2.add_run(); r4.text = "0.0171    "; r4.font.bold = True; r4.font.size = Pt(14); r4.font.color.rgb = C_WHITE
    r5 = p2.add_run(); r5.text = "Precision: "; r5.font.bold = True; r5.font.size = Pt(13); r5.font.color.rgb = C_AMBER
    r6 = p2.add_run(); r6.text = "0.2317    "; r6.font.bold = True; r6.font.size = Pt(14); r6.font.color.rgb = C_WHITE
    r7 = p2.add_run(); r7.text = "Recall: "; r7.font.bold = True; r7.font.size = Pt(13); r7.font.color.rgb = C_AMBER
    r8 = p2.add_run(); r8.text = "0.0089    "; r8.font.bold = True; r8.font.size = Pt(14); r8.font.color.rgb = C_WHITE
    r9 = p2.add_run(); r9.text = "Water Pixels: "; r9.font.bold = True; r9.font.size = Pt(11); r9.font.color.rgb = C_MUTED
    r10 = p2.add_run(); r10.text = "4,277"; r10.font.bold = True; r10.font.size = Pt(11); r10.font.color.rgb = C_TEXT_LIGHT

    create_card(s14, Inches(0.6), Inches(2.85), Inches(12.133), Inches(3.95), bg_color=C_CARD_DARK)
    img_s14 = os.path.join(assets_dir, "outputs/visualizations/final_test/3_difficult_prediction_Sri-Lanka_450918.png")
    if os.path.exists(img_s14):
        s14.shapes.add_picture(img_s14, Inches(0.8), Inches(3.0), width=Inches(11.73))
        tb_cap = s14.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.73), Inches(0.9))
        tf = tb_cap.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "Panels (L to R): 1. Input SAR False-Color (VV, VH, Ratio) | 2. Ground Truth Mask | 3. Predicted Flood Mask | 4. Model Overlay | 5. Flood Probability Map"
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_AMBER
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = "Physical Failure Mode: Dense tropical vegetation canopy and submerged tree trunks create corner-reflector 'double-bounce' backscatter, masking standing water beneath the foliage."
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s14, """
Slide 14 presents an honest, transparent analysis of our most challenging test case: Sri-Lanka_450918.
The model obtained an IoU of 0.0086 and Dice of 0.0171.
We explicitly include this failure mode to demonstrate the physical constraints of single-pass C-band SAR:
In dense tropical rainforests, floodwaters under the tree canopy do not exhibit specular reflection.
Instead, the water surface and vertical tree trunks form 90-degree corner reflectors, bouncing microwave energy back to the satellite antenna in a phenomenon known as 'double-bounce'.
As a result, the flooded forest appears bright rather than dark, causing the model to miss submerged vegetation.
This provides a clear justification for future multi-sensor and bi-temporal SAR enhancements.
""")

    # =========================================================================
    # SLIDE 15: FLOOD MAPPING & NOMINAL FLOODED AREA (LARGE SOMALIA VISUAL)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15)
    add_header(s15, "Geospatial Analysis", "Flood Mapping & Nominal Flooded Area Estimation", "Transforming continuous probability heatmaps into actionable disaster management metrics")
    add_footer(s15, 15)

    # Top: Geospatial Explanation Card (12.13 inches wide, 1.4 inches high)
    create_card(s15, Inches(0.6), Inches(1.75), Inches(12.133), Inches(1.4))
    add_card_header(s15, Inches(0.8), Inches(1.85), Inches(11.7), "Geospatial Flow & Nominal Area Terminology", color=C_CYAN, size=12.5)

    tb_geo = s15.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.9))
    tf = tb_geo.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "• Continuous Probability Map: Sigmoid output represents pixel-level flood probability p̂ ∈ [0, 1]. Termed 'Flood Probability Map' rather than confidence map.\n• Thresholding & Nominal Area: Pixels with p̂ ≥ τ (default τ = 0.50) are aggregated: Nominal Flooded Area ≈ N_flood × 100 m² / 10⁶ = km².\n• WGS84 Scaling Note: Because unprojected EPSG:4326 pixel width scales with cos(latitude), area is explicitly reported as a nominal equatorial approximation."
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_LIGHT

    # Bottom: Large Somalia Regional Flood Image (12.13 inches wide, 3.5 inches high)
    create_card(s15, Inches(0.6), Inches(3.3), Inches(12.133), Inches(3.5), bg_color=C_CARD_DARK)
    add_card_header(s15, Inches(0.8), Inches(3.45), Inches(11.7), "Somalia Regional Flood Inundation Case Study", color=C_GREEN, size=12.5)

    img_s15 = os.path.join(assets_dir, "outputs/visualizations/predictions/Somalia_989553_regional_prediction.png")
    if os.path.exists(img_s15):
        s15.shapes.add_picture(img_s15, Inches(0.8), Inches(3.85), width=Inches(11.73))
        tb_cap = s15.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.73), Inches(0.3))
        tf = tb_cap.text_frame
        p = tf.paragraphs[0]
        p.text = "Regional Mapping Output: Input SAR Composite, Ground Truth Target, Predicted Flood Mask, and Color Overlay"
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

    add_speaker_notes(s15, """
Slide 15 explains how model probabilities translate into geospatial flood maps.
Our network outputs continuous values between 0 and 1, which we term the 'Flood Probability Map'.
Applying an adjustable decision threshold (defaulting to 0.50) generates the binary flood polygon.
We then compute the 'Nominal Flooded Area' in square kilometers based on nominal 10-meter pixel resolution (100 square meters per pixel).
We explicitly clarify that because Sen1Floods11 rasters are unprojected WGS84 (EPSG:4326), pixel width varies with the cosine of latitude. Therefore, we report 'Nominal Flooded Area' rather than claiming exact geodesic area.
The large visual below displays the regional inundation mapping for Somalia.
""")

    # =========================================================================
    # SLIDE 16: STREAMLIT APPLICATION (LARGE UI VISUALS + LOCAL DEMO)
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_bg(s16)
    add_header(s16, "Interactive Demonstration", "Interactive Streamlit Web Dashboard (`app.py`)", "Dual-mode web interface supporting benchmark exploration and custom GeoTIFF upload")
    add_footer(s16, 16)

    # Top summary cards (2 columns)
    w_half = Inches(5.9)
    create_card(s16, Inches(0.6), Inches(1.75), w_half, Inches(1.5))
    add_card_header(s16, Inches(0.8), Inches(1.85), w_half - Inches(0.4), "Mode A: Benchmark Test Explorer", color=C_CYAN, size=12)
    tb_a = s16.shapes.add_textbox(Inches(0.8), Inches(2.2), w_half - Inches(0.4), Inches(0.95))
    tf = tb_a.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "• Dropdown selector for all 90 official test chips.\n• Live 5-panel diagnostics: SAR composite, ground truth, prediction, overlay, and probability heatmap.\n• Instant quantitative metrics (IoU, Dice, Precision, Recall)."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_TEXT_LIGHT

    create_card(s16, Inches(6.8), Inches(1.75), w_half + Inches(0.033), Inches(1.5))
    add_card_header(s16, Inches(7.0), Inches(1.85), w_half - Inches(0.4), "Mode B: Custom GeoTIFF Upload", color=C_GREEN, size=12)
    tb_b = s16.shapes.add_textbox(Inches(7.0), Inches(2.2), w_half - Inches(0.4), Inches(0.95))
    tf = tb_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "• Ingests arbitrary dual-polarization Sentinel-1 GeoTIFFs.\n• Validates channel counts (≥ 2 bands) and normalizes dB.\n• Interactive threshold slider (0.10 to 0.90) with sub-50ms inference."
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_TEXT_LIGHT

    # Bottom: Large Dashboard Output Visual & Demo Callout
    create_card(s16, Inches(0.6), Inches(3.4), Inches(12.133), Inches(3.4), bg_color=C_CARD_DARK)
    add_card_header(s16, Inches(0.8), Inches(3.55), Inches(8.0), "Dashboard Diagnostic Visualizer Output", color=C_WHITE, size=12)

    # Local Demo Callout Box
    demo_btn = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(3.5), Inches(4.0), Inches(0.4))
    demo_btn.fill.solid()
    demo_btn.fill.fore_color.rgb = C_BORDER_CYAN
    demo_btn.line.fill.background()
    tf = demo_btn.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "LOCAL DEMO → http://localhost:8501"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    img_s16 = os.path.join(assets_dir, "outputs/visualizations/predictions/USA_994009_regional_prediction.png")
    if os.path.exists(img_s16):
        s16.shapes.add_picture(img_s16, Inches(0.8), Inches(4.0), width=Inches(11.73))

    add_speaker_notes(s16, """
Slide 16 presents our interactive Streamlit demonstration dashboard, implemented in app.py.
Mode A is the Benchmark Test Explorer, allowing examiners to inspect any of the 90 official test chips with full 5-panel visual diagnostics and live metric calculations.
Mode B supports custom GeoTIFF uploads for testing arbitrary Sentinel-1 files.
Model inference takes under 50 milliseconds thanks to weight caching with @st.cache_resource.
We are ready to switch to the live local demo at localhost:8501 during questions.
""")

    # =========================================================================
    # SLIDE 17: LIMITATIONS & FUTURE SCOPE
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_bg(s17)
    add_header(s17, "Critical Analysis & Roadmap", "Academic Limitations & Future Scope", "Objective evaluation of current constraints and proposed engineering enhancements")
    add_footer(s17, 17)

    col_w = Inches(5.9)
    card_h = Inches(5.0)

    # Left: Limitations
    create_card(s17, Inches(0.6), Inches(1.75), col_w, card_h)
    add_card_header(s17, Inches(0.8), Inches(1.9), col_w - Inches(0.4), "Current Academic Limitations", color=C_AMBER, size=13)

    tb_lim = s17.shapes.add_textbox(Inches(0.8), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb_lim.text_frame
    tf.word_wrap = True
    limits = [
        ("• Specular Surface False Positives: ", "Smooth dry surfaces such as airport runways and flat playas specularly reflect radar pulses, causing occasional false alarms."),
        ("• Canopy Volume False Negatives: ", "Flooded vegetation and dense tree canopies cause double-bounce backscatter, obscuring standing water beneath foliage."),
        ("• Unlabeled No-Data Pixels: ", "Sen1Floods11 chips contain -1 invalid pixels that must be masked out during training."),
        ("• Historical JRC Baseline: ", "Sen1Floods11 relies on the 30-year JRC permanence raster rather than paired pre-flood SAR acquisitions."),
        ("• Nominal Equatorial Area: ", "WGS84 angular resolution requires nominal metric assumptions rather than local UTM projection.")
    ]
    for h, b in limits:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_AMBER
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(4)

    # Right: Future Scope
    create_card(s17, Inches(6.8), Inches(1.75), col_w + Inches(0.033), card_h, border_color=C_BORDER_CYAN)
    add_card_header(s17, Inches(7.0), Inches(1.9), col_w - Inches(0.4), "Future Research & Engineering Scope", color=C_CYAN, size=13)

    tb_fut = s17.shapes.add_textbox(Inches(7.0), Inches(2.35), col_w - Inches(0.4), Inches(4.2))
    tf = tb_fut.text_frame
    tf.word_wrap = True
    futures = [
        ("• Bi-Temporal SAR Difference Networks: ", "Ingesting co-registered pre-flood and post-flood SAR pairs to directly isolate floodwater from permanent water bodies."),
        ("• Multi-Sensor & DEM Fusion: ", "Fusing Sentinel-1 SAR with Copernicus DEM elevation rasters and Sentinel-2 optical data to enforce topographical gravity constraints."),
        ("• Infrastructure Damage Intersection: ", "Overlaying OpenStreetMap building footprints, hospitals, and road networks onto predicted flood masks for evacuation planning."),
        ("• Cloud-Native Automated Alerts: ", "Deploying as a serverless container hooked to Copernicus Open Access Hub webhooks for automated near-real-time flood alerting.")
    ]
    for h, b in futures:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_CYAN_LIGHT
        r2 = p.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_after = Pt(6)

    add_speaker_notes(s17, """
Slide 17 provides an objective assessment of limitations and future scope.
Current limitations:
1. False positives on smooth dry land like airport tarmac due to specular radar reflection.
2. False negatives in dense forests where double-bounce backscatter masks standing water.
3. Reliance on the JRC 30-year permanence raster rather than paired pre-flood SAR acquisitions.
4. Nominal equatorial area estimation due to unprojected WGS84 coordinates.
Future work includes:
1. Bi-temporal SAR difference networks with pre-and-post flood SAR pairs.
2. Integrating digital elevation models (DEM) to filter elevated terrain.
3. Intersecting flood polygons with OpenStreetMap road and building vector layers to identify stranded communities.
""")

    # =========================================================================
    # SLIDE 18: CONCLUSION
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    set_bg(s18)
    create_card(s18, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3), bg_color=C_CARD_DARK, border_color=C_GREEN, line_width=1.5)

    badge = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.9), Inches(3.2), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_GREEN
    badge.line.fill.background()
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "CONCLUSION & VIVA DEFENSE"
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_BG
    p.alignment = PP_ALIGN.CENTER

    tb_title = s18.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(11.0), Inches(0.6))
    tf = tb_title.text_frame
    p = tf.paragraphs[0]
    p.text = "From Satellite Radar to Verified Flood Maps"
    p.font.name = FONT_HEADING
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    box_w = Inches(2.6)
    box_h = Inches(2.2)
    achievements = [
        ("ROBUST PIPELINE", "Sentinel-1 VV/VH", "Automated decibel normalization & strided subsampling pipeline with 0 data leakage.", C_CYAN),
        ("OPTIMIZED U-NET", "7.76M Parameters", "Group Normalization (G=8) ensured stable convergence with batch size 2.", C_GREEN),
        ("TEST BENCHMARK", "0.5548 Test IoU", "Evaluated on 100% of 90 test chips (20.5M pixels) with 0.7137 Dice & 93.34% Acc.", C_CYAN_LIGHT),
        ("VERIFIED QUALITY", "39/39 Tests Passed", "Full test suite passing, live Streamlit UI & sub-50ms inference latency.", C_WHITE)
    ]
    for i, (head, stat, desc, col) in enumerate(achievements):
        x = Inches(1.0 + i * 2.8)
        create_card(s18, x, Inches(2.1), box_w, box_h, bg_color=C_CARD, border_color=C_BORDER)

        tb = s18.shapes.add_textbox(x + Inches(0.1), Inches(2.2), box_w - Inches(0.2), box_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_MUTED
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = stat
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)
        p2.space_after = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(9)
        p3.font.color.rgb = C_TEXT_LIGHT
        p3.alignment = PP_ALIGN.CENTER

    tb_ty = s18.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.333), Inches(1.8))
    tf = tb_ty.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Thank You for Your Time & Consideration"
    p.font.name = FONT_HEADING
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_CYAN
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Open for Questions, Evaluation, and Live Demonstration"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)

    p3 = tf.add_paragraph()
    p3.text = "Project Repository: AI-Based Flood Detection and Mapping  |  Sen1Floods11 • Sentinel-1 SAR • PyTorch U-Net"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(10)
    p3.font.color.rgb = C_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(10)

    add_speaker_notes(s18, """
In summary, we have built and audited an end-to-end deep learning semantic segmentation system for satellite radar flood mapping.
By leveraging Sentinel-1 dual polarization, Group Normalization (G=8), and a masked BCE-Dice loss, our model achieves 0.5548 Test IoU and 0.7137 Dice across all 90 chips of the official benchmark without data leakage.
Our complete test suite of 39 unit tests passes with 100% success, and our Streamlit application provides real-time flood mapping and GeoTIFF upload.
Thank you for your attention. We are ready for your questions and the live demonstration.
""")

    # Save filenames
    out_final_upper = os.path.join(assets_dir, "AI_Flood_Detection_Minor_Project_FINAL.pptx")
    out_final = os.path.join(assets_dir, "AI_Flood_Detection_Minor_Project_Final.pptx")
    out_orig = os.path.join(assets_dir, "AI_Flood_Detection_Minor_Project.pptx")
    prs.save(out_final_upper)
    prs.save(out_final)
    prs.save(out_orig)
    print(f"Successfully generated {out_final_upper}, {out_final}, and {out_orig} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    build_presentation()
